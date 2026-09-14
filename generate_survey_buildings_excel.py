import os
import sys
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
import pandas as pd
import numpy as np
import unicodedata

def generate_report(survey_id='KS003'):
    survey_dir = f'data/output/{survey_id}'
    excel_path = os.path.join(survey_dir, f'{survey_id}.xlsx')
    output_excel = os.path.join(survey_dir, f'Bao_Cao_Phan_Tich_Quy_Hoach_{survey_id}.xlsx')
    config_path = os.path.join(survey_dir, 'khao_sat_config.json')
    
    grid_m = 4
    total_points = 82824
    if os.path.exists(config_path):
        try:
            with open(config_path, 'r', encoding='utf-8') as f:
                cfg = json.load(f)
                grid_m = cfg.get('grid', 4)
                total_points = cfg.get('total_points', 82824)
        except Exception as e:
            print(f"Notice: Config read error {e}")

    print(f"=== Generating Comprehensive Report for {survey_id} ===")
    print(f"1. Loading raw data from: {excel_path}")
    if not os.path.exists(excel_path):
        print(f"Error: {excel_path} not found!")
        return

    df1 = pd.read_excel(excel_path, sheet_name=0, header=3)
    df2 = pd.read_excel(excel_path, sheet_name=1, header=3)
    df3 = pd.read_excel(excel_path, sheet_name=2, header=3)
    df4 = pd.read_excel(excel_path, sheet_name=3, header=3)

    def normalize(s):
        if not isinstance(s, str):
            return ''
        return unicodedata.normalize('NFC', s).strip()

    df2['norm_land'] = df2['Chức Năng Sử Dụng Đất'].apply(normalize)
    df2['norm_detail'] = df2['Chức Năng Chi Tiết (SQHKT)'].apply(normalize)
    df2['norm_block'] = df2['Mã Ô Phố'].apply(lambda x: '' if str(x).strip() in ['-', '', 'None', 'nan'] else f"[{str(x).strip()}] ")

    def build_label(r):
        code = r['norm_block']
        name = r['norm_land']
        dt = r['Diện Tích Ô (m²)']
        pct = r['Tỷ Lệ Chiếm Thửa (%)'] * 100
        return f"{code}{name}: {dt:,.2f} m² ({pct:.1f}%)"

    df2['Nhãn Ô Chi Tiết'] = df2.apply(build_label, axis=1)

    def is_housing(name):
        n = name.lower()
        return any(k in n for k in ['đất ở', 'dân cư', 'nhà ở', 'nhóm nhà ở'])

    def is_facility(name):
        n = name.lower()
        return any(k in n for k in [
            'thương mại', 'tmdv', 'sxkd', 'cơ quan', 'phức hợp', 'hỗn hợp',
            'công trình', 'công cộng', 'giáo dục', 'trường học', 'y tế', 'hành chính', 'hành chánh',
            'tôn giáo', 'di tích', 'quảng trường', 'khách sạn', 'ga depot', 'công nghiệp', 'tiểu thủ', 'trạm cấp nước', 'kho tàng'
        ]) and not is_housing(name)

    def get_facility_type(name):
        n = name.lower()
        if any(k in n for k in ['thương mại', 'tmdv']):
            return 'Thương Mại - Dịch Vụ'
        if any(k in n for k in ['phức hợp', 'hỗn hợp']):
            return 'Phức Hợp / Hỗn Hợp'
        if any(k in n for k in ['khách sạn']):
            return 'Khách Sạn & Dịch Vụ'
        if any(k in n for k in ['cơ quan', 'sxkd', 'hành chính', 'hành chánh']):
            return 'Trụ Sở Cơ Quan / Hành Chính'
        if any(k in n for k in ['giáo dục', 'trường học']):
            return 'Công Trình Giáo Dục'
        if any(k in n for k in ['y tế']):
            return 'Công Trình Y Tế / Bệnh Viện'
        if any(k in n for k in ['tôn giáo', 'di tích']):
            return 'Tôn Giáo / Di Tích'
        if any(k in n for k in ['ga depot', 'depot']):
            return 'Ga Depot Metro'
        if any(k in n for k in ['công nghiệp', 'tiểu thủ', 'kho tàng', 'trạm cấp nước']):
            return 'Công Nghiệp & Hạ Tầng'
        if any(k in n for k in ['công trình', 'công cộng', 'quảng trường']):
            return 'Công Trình Công Cộng'
        return 'Khác'

    def get_category_broad(name):
        n = name.lower()
        if is_housing(name):
            return '1. Đất Ở / Nhà Ở & Dân Cư'
        if any(k in n for k in ['phức hợp', 'hỗn hợp']):
            return '2. Đất Phức Hợp / Hỗn Hợp'
        if any(k in n for k in ['thương mại', 'tmdv', 'sxkd', 'cơ quan']):
            return '3. Thương Mại - Dịch Vụ & Cơ Quan'
        if any(k in n for k in ['công cộng', 'giáo dục', 'trường học', 'y tế', 'hành chính', 'hành chánh', 'tôn giáo', 'di tích', 'quảng trường']):
            return '4. Công Trình Công Cộng & An Sinh'
        if any(k in n for k in ['giao thông', 'ga depot', 'depot', 'mặt nước', 'kênh', 'rạch', 'sông', 'hạ tầng', 'cấp nước', 'kho tàng']):
            return '6. Giao Thông & Hạ Tầng Kỹ Thuật'
        if any(k in n for k in ['công nghiệp', 'tiểu thủ', 'quân sự', 'quốc phòng']):
            return '7. Công Nghiệp & Quốc Phòng'
        if any(k in n for k in ['cây xanh', 'công viên', 'cv -', 'cv-', 'thể dục thể thao']):
            return '5. Cây Xanh, Công Viên & TDTT'
        return '8. Đất Khác'

    df2['Nhóm Quy Hoạch Tham Khảo'] = df2['norm_land'].apply(get_category_broad)
    df2['Loại Công Trình'] = df2['norm_land'].apply(lambda x: get_facility_type(x) if is_facility(x) else '-')

    tot_survey_area = df1['Diện Tích Thửa (m²)'].sum()

    # Aggregations for raw catalog
    raw_types_df = df2.groupby('norm_land').agg(
        Nhóm_Tham_Khảo=('Nhóm Quy Hoạch Tham Khảo', 'first'),
        Số_Ô=('STT Thửa', 'count'),
        Số_Thửa_Tiếp_Giáp=('STT Thửa', 'nunique'),
        Tổng_DT_Ô_m2=('Diện Tích Ô (m²)', 'sum'),
        DT_Trung_Bình_Ô_m2=('Diện Tích Ô (m²)', 'mean'),
        Tầng_Cao_Max=('Tầng Cao Cho Phép (tầng)', lambda x: pd.to_numeric(x, errors='coerce').max()),
        Mật_Độ_Max=('Mật Độ XD Tối Đa (%)', lambda x: pd.to_numeric(x, errors='coerce').max()),
        HSSDĐ_Max=('Hệ Số Sử Dụng Đất (HSSDĐ)', lambda x: pd.to_numeric(x, errors='coerce').max()),
        Mã_Ô_Phố_Tiêu_Biểu=('Mã Ô Phố', lambda x: ', '.join([str(i) for i in x.unique() if str(i).strip() not in ['-', '', 'None', 'nan']][:6]))
    ).reset_index().rename(columns={'norm_land': 'Tên Loại Đất Thô (SQHKT)'})

    raw_types_df['Tỷ Lệ DT Trên Tổng Khảo Sát (%)'] = raw_types_df['Tổng_DT_Ô_m2'] / tot_survey_area * 100
    raw_types_df = raw_types_df.sort_values(by='Tổng_DT_Ô_m2', ascending=False).reset_index(drop=True)
    raw_types_df.insert(0, 'STT', range(1, len(raw_types_df) + 1))

    # Per parcel aggregations
    # 1. Housing
    housing_df2 = df2[df2['norm_land'].apply(is_housing)]
    housing_agg = housing_df2.groupby('STT Thửa').agg(
        DT_Nha_O=('Diện Tích Ô (m²)', 'sum'),
        Ty_Le_Nha_O=('Tỷ Lệ Chiếm Thửa (%)', 'sum'),
        Loai_Nha_O=('norm_land', lambda x: ', '.join(x.unique())),
        Tang_Cao_Nha_O=('Tầng Cao Cho Phép (tầng)', lambda x: pd.to_numeric(x, errors='coerce').max()),
        Mat_Do_Nha_O=('Mật Độ XD Tối Đa (%)', lambda x: pd.to_numeric(x, errors='coerce').max()),
        HSSDD_Nha_O=('Hệ Số Sử Dụng Đất (HSSDĐ)', lambda x: pd.to_numeric(x, errors='coerce').max())
    ).reset_index()
    housing_dict = {r['STT Thửa']: r for _, r in housing_agg.iterrows()}

    # 2. Facilities
    facility_df2 = df2[df2['norm_land'].apply(is_facility)]
    facility_agg = facility_df2.groupby('STT Thửa').agg(
        DT_Cong_Trinh=('Diện Tích Ô (m²)', 'sum'),
        Ty_Le_Cong_Trinh=('Tỷ Lệ Chiếm Thửa (%)', 'sum'),
        Loai_Cong_Trinh=('Loại Công Trình', lambda x: ', '.join([i for i in x.unique() if i != '-'])),
        Ten_Dat_Cong_Trinh=('norm_land', lambda x: ', '.join(x.unique())),
        Tang_Cao_CT=('Tầng Cao Cho Phép (tầng)', lambda x: pd.to_numeric(x, errors='coerce').max()),
        Mat_Do_CT=('Mật Độ XD Tối Đa (%)', lambda x: pd.to_numeric(x, errors='coerce').max()),
        HSSDD_CT=('Hệ Số Sử Dụng Đất (HSSDĐ)', lambda x: pd.to_numeric(x, errors='coerce').max())
    ).reset_index()
    facility_dict = {r['STT Thửa']: r for _, r in facility_agg.iterrows()}

    # 3. Traffic
    traffic_agg = df2[df2['norm_land'] == 'Đất giao thông'].groupby('STT Thửa').agg(
        DT_Giao_Thong=('Diện Tích Ô (m²)', 'sum'),
        Ty_Le_Giao_Thong=('Tỷ Lệ Chiếm Thửa (%)', 'sum')
    ).reset_index()
    traffic_dict = {r['STT Thửa']: r for _, r in traffic_agg.iterrows()}

    parcels_with_1500 = set(df4[df4['Loại Quy Hoạch'] == 'Quy hoạch chi tiết 1/500']['STT Thửa'].unique())
    parcels_with_dccb = set(df4[df4['Loại Quy Hoạch'] == 'Điều chỉnh cục bộ (DCCB)']['STT Thửa'].unique())

    def enrich_parcel(row):
        stt = row['STT']
        h = housing_dict.get(stt, None)
        f = facility_dict.get(stt, None)
        t = traffic_dict.get(stt, None)
        
        dt_o = round(h['DT_Nha_O'], 2) if h is not None else 0.0
        pct_o = round(h['Ty_Le_Nha_O'] * 100, 2) if h is not None else 0.0
        loai_o = h['Loai_Nha_O'] if h is not None else 'Không'
        tc_o = h['Tang_Cao_Nha_O'] if h is not None else np.nan
        md_o = h['Mat_Do_Nha_O'] if h is not None else np.nan
        hs_o = h['HSSDD_Nha_O'] if h is not None else np.nan
        
        dt_ct = round(f['DT_Cong_Trinh'], 2) if f is not None else 0.0
        pct_ct = round(f['Ty_Le_Cong_Trinh'] * 100, 2) if f is not None else 0.0
        phan_loai_ct = f['Loai_Cong_Trinh'] if f is not None else 'Không'
        ten_dat_ct = f['Ten_Dat_Cong_Trinh'] if f is not None else 'Không'
        tc_ct = f['Tang_Cao_CT'] if f is not None else np.nan
        md_ct = f['Mat_Do_CT'] if f is not None else np.nan
        hs_ct = f['HSSDD_CT'] if f is not None else np.nan
        
        dt_gt = round(t['DT_Giao_Thong'], 2) if t is not None else 0.0
        pct_gt = round(t['Ty_Le_Giao_Thong'] * 100, 2) if t is not None else 0.0
        
        status_1500 = 'Có QH 1/500' if stt in parcels_with_1500 else ('Có ĐCCB' if stt in parcels_with_dccb else 'Không')
        
        if dt_o > 0 and dt_ct > 0:
            ht = '🏡🏢 Hỗn Hợp Nhà Ở & Công Trình'
        elif dt_o > 0:
            ht = '🏡 Đất Có Nhà Ở / Dân Cư'
        elif dt_ct > 0:
            ht = '🏢 Đất Có Công Trình Xây Dựng'
        elif dt_gt > 0:
            ht = '🛣️ Đất Lộ Giới Giao Thông Thuần'
        else:
            ht = '⚪ Đất Khác / Chuyên Dùng'
            
        return pd.Series({
            'Phân Loại Thực Địa': ht,
            'Có Nhà Ở': 'Có' if dt_o > 0 else 'Không',
            'DT Nhà Ở (m²)': dt_o,
            'Tỷ Lệ Nhà Ở (%)': pct_o,
            'Loại Đất Nhà Ở Thô': loai_o,
            'Tầng Cao Nhà Ở': tc_o,
            'Mật Độ Nhà Ở (%)': md_o,
            'HSSDĐ Nhà Ở': hs_o,
            'Có Công Trình XD': 'Có' if dt_ct > 0 else 'Không',
            'DT Công Trình (m²)': dt_ct,
            'Tỷ Lệ Công Trình (%)': pct_ct,
            'Phân Loại Công Trình': phan_loai_ct,
            'Tên Đất Công Trình Thô': ten_dat_ct,
            'Tầng Cao CT': tc_ct,
            'Mật Độ CT (%)': md_ct,
            'HSSDĐ CT': hs_ct,
            'DT Giao Thông Lộ Giới (m²)': dt_gt,
            'Tỷ Lệ Giao Thông (%)': pct_gt,
            'Dự Án 1/500 & ĐCCB': status_1500
        })

    df1_enriched = pd.concat([df1, df1.apply(enrich_parcel, axis=1)], axis=1)

    tot_housing_parcels = (df1_enriched['Có Nhà Ở'] == 'Có').sum()
    tot_facility_parcels = (df1_enriched['Có Công Trình XD'] == 'Có').sum()
    tot_built_parcels = ((df1_enriched['Có Nhà Ở'] == 'Có') | (df1_enriched['Có Công Trình XD'] == 'Có')).sum()
    tot_traffic_parcels = (df1_enriched['DT Giao Thông Lộ Giới (m²)'] > 0).sum()

    print("Enrichment complete.")
    print(f"Total Parcels: {len(df1_enriched):,}")
    print(f"Parcels with Housing: {tot_housing_parcels:,} ({tot_housing_parcels/len(df1)*100:.2f}%)")
    print(f"Parcels with Facilities: {tot_facility_parcels:,} ({tot_facility_parcels/len(df1)*100:.2f}%)")
    print(f"Parcels with Housing OR Facilities: {tot_built_parcels:,} ({tot_built_parcels/len(df1)*100:.2f}%)")

    # Workbook Creation
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    FONT_NAME = 'Arial'
    font_title = Font(name=FONT_NAME, size=13, bold=True, color='FFFFFF')
    font_subtitle = Font(name=FONT_NAME, size=9, italic=True, color='E2E8F0')
    font_section = Font(name=FONT_NAME, size=11, bold=True, color='1E3A8A')
    font_table_hdr = Font(name=FONT_NAME, size=9, bold=True, color='FFFFFF')
    font_kpi_num = Font(name=FONT_NAME, size=15, bold=True, color='1E3A8A')
    font_kpi_label = Font(name=FONT_NAME, size=9, bold=True, color='475569')
    font_kpi_sub = Font(name=FONT_NAME, size=8, italic=True, color='64748B')
    font_cell = Font(name=FONT_NAME, size=9)
    font_cell_bold = Font(name=FONT_NAME, size=9, bold=True)

    fill_navy_hdr = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    fill_blue_hdr = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    fill_teal_hdr = PatternFill(start_color='0D9488', end_color='0D9488', fill_type='solid')
    fill_purple_hdr = PatternFill(start_color='4338CA', end_color='4338CA', fill_type='solid')
    fill_amber_hdr = PatternFill(start_color='B45309', end_color='B45309', fill_type='solid')
    fill_kpi_bg = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
    fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    fill_highlight_green = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')
    fill_highlight_amber = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
    fill_highlight_blue = PatternFill(start_color='DBEAFE', end_color='DBEAFE', fill_type='solid')
    fill_total_row = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')

    thin_side = Side(border_style='thin', color='CBD5E1')
    thick_bottom = Side(border_style='medium', color='1E3A8A')
    double_bottom = Side(border_style='double', color='1E3A8A')

    border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    border_header = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thick_bottom)
    border_total = Border(left=thin_side, right=thin_side, top=thin_side, bottom=double_bottom)

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center')

    # Header explanations & tooltips
    HEADER_TOOLTIPS = {
        'Tầng Cao Max': 'Số tầng cao tối đa được phép xây dựng công trình theo đồ án quy hoạch phân khu 1/2000 được phê duyệt.',
        'Tầng Cao Cho Phép': 'Số tầng cao tối đa được phép xây dựng công trình theo đồ án quy hoạch phân khu 1/2000.',
        'Tầng Cao Cho Phép (tầng)': 'Số tầng cao tối đa được phép xây dựng công trình theo đồ án quy hoạch phân khu 1/2000.',
        'Tầng Cao Nhà Ở': 'Tầng cao tối đa cho phép xây dựng đối với phần diện tích đất ở/nhà ở trên thửa.',
        'Tầng Cao CT': 'Tầng cao tối đa cho phép xây dựng đối với phần diện tích công trình trên thửa.',
        'Mật Độ XD Max (%)': 'Mật độ xây dựng thuần tối đa = (Diện tích xây dựng chiếm đất / Tổng diện tích lô đất) * 100%. Ví dụ: Mật độ 60% trên lô 100m² thì diện tích sàn trệt tối đa là 60m², còn lại làm sân lùi/cây xanh.',
        'Mật Độ XD (%)': 'Mật độ xây dựng thuần tối đa (%) cho phép xây dựng công trình.',
        'Mật Độ Nhà Ở (%)': 'Mật độ xây dựng tối đa (%) cho phép đối với phần đất ở/nhà ở.',
        'Mật Độ CT (%)': 'Mật độ xây dựng tối đa (%) cho phép đối với phần đất công trình.',
        'HSSDĐ Max': 'Hệ số sử dụng đất tối đa (FAR - Floor Area Ratio) = Tổng diện tích sàn xây dựng tất cả các tầng (trừ hầm/kỹ thuật) / Tổng diện tích lô đất. Ví dụ: Lô 100m² có HSSDĐ 3.5 thì tổng sàn xây dựng tối đa là 350m².',
        'HSSDĐ': 'Hệ số sử dụng đất tối đa (FAR) = Tổng diện tích sàn xây dựng / Diện tích lô đất.',
        'HSSDĐ Nhà Ở': 'Hệ số sử dụng đất tối đa cho phép xây dựng đối với phần đất ở.',
        'HSSDĐ CT': 'Hệ số sử dụng đất tối đa cho phép xây dựng đối với phần đất công trình.',
        'Dự Án 1/500 & ĐCCB': 'Tình trạng pháp lý quy hoạch: Cho biết thửa đất có nằm trong ranh đồ án Quy hoạch chi tiết 1/500 (dự án phát triển cụ thể) hoặc có Quyết định Điều chỉnh cục bộ (ĐCCB) hay không.',
        'Lộ Giới Tiếp Giáp': 'Bề rộng lộ giới (chỉ giới đường đỏ) của tuyến đường tiếp giáp thửa đất. Xác định khoảng lùi xây dựng và phần diện tích đất sẽ bị thu hồi/trừ đi khi mở đường theo quy hoạch giao thông.',
        'STT': 'Số thứ tự của dòng dữ liệu trong bảng.',
        'STT Thửa': f'Số thứ tự định danh thửa đất khảo sát (từ 1 đến {len(df1):,}).',
        'STT Ô': f'Số thứ tự định danh ô quy hoạch phân khu chức năng (từ 1 đến {len(df2):,}).',
        'Mã Thửa Đất': 'Mã định danh duy nhất của thửa đất theo hệ thống địa chính (kết hợp Quận-Phường-Tờ-Thửa).',
        'Số Tờ': 'Số hiệu tờ bản đồ địa chính chính quy nơi thửa đất tọa lạc.',
        'Số Thửa': 'Số hiệu thửa đất trên tờ bản đồ địa chính tương ứng.',
        'Quận / Huyện': 'Đơn vị hành chính cấp Quận/Huyện quản lý thửa đất.',
        'Phường / Xã': 'Đơn vị hành chính cấp Phường/Xã nơi thửa đất tọa lạc.',
        'Mã Ô Phố': 'Mã ký hiệu định danh ô quy hoạch phân khu trong bản đồ 1/2000 (Ví dụ: [II.26], [III.9], [I/89]...).',
        'Toạ Độ Khảo Sát': 'Tọa độ tâm thửa đất (Kinh độ, Vĩ độ theo hệ WGS84), dùng để định vị chính xác trên Google Maps và GIS.',
        'Tên Loại Đất Thô (SQHKT)': f'Tên danh mục chức năng sử dụng đất gốc từ hệ thống Sở Quy Hoạch - Kiến Trúc TP.HCM ({len(raw_types_df)} loại thô, chưa gộp).',
        'Chức Năng Sử Dụng Đất (Thô)': 'Chức năng sử dụng đất phân khu 1/2000 nguyên bản theo dữ liệu SQHKT.',
        'Chức Năng Chi Tiết (SQHKT)': 'Mô tả chi tiết mục đích sử dụng đất và chức năng quy hoạch theo hồ sơ pháp lý SQHKT.',
        'Nhóm Tham Khảo': 'Nhóm phân loại chức năng mở rộng (Đất ở, TMDV, Công trình công cộng, Giao thông...) để dễ phân nhóm tổng quan.',
        'Nhãn Ô Chi Tiết Đầy Đủ': 'Nhãn gộp trực quan theo chuẩn: [Mã Ô] Tên loại đất: Diện tích m² (Tỷ lệ %).',
        'Phân Loại Thực Địa': 'Phân nhóm phục vụ khảo sát hiện trường: Nhà ở dân cư, Công trình xây dựng, Hỗn hợp hoặc Lộ giới giao thông.',
        'Phân Loại Công Trình': 'Nhóm công trình xây dựng (TMDV, Phức hợp, Cơ quan, Giáo dục, Y tế, Khách sạn, Ga Depot, Công nghiệp...).',
        'Loại Đất Nhà Ở Thô': 'Tên chức năng đất ở cụ thể (Đất ở hiện hữu, Đất ở cải tạo chỉnh trang, Đất dân cư dự kiến...).',
        'Tên Đất Công Trình Thô': 'Tên chức năng đất công trình cụ thể theo dữ liệu thô SQHKT.',
        'Cơ Cấu Quy Hoạch Thô (SQHKT)': 'Tổng hợp toàn bộ các ô quy hoạch thành phần cấu thành nên thửa đất kèm diện tích và tỷ lệ %.',
        'Tổng DT Thửa (m²)': 'Tổng diện tích toàn bộ khuôn viên thửa đất theo ranh địa chính (m²).',
        'Diện Tích Thửa (m²)': 'Tổng diện tích toàn bộ khuôn viên thửa đất theo ranh địa chính (m²).',
        'Tổng Diện Tích (m²)': 'Tổng diện tích các thửa đất (m²).',
        'Tổng DT (ha)': 'Tổng diện tích quy đổi ra đơn vị Hecta (1 ha = 10.000 m²).',
        'Tổng Diện Tích Ô (m²)': 'Tổng diện tích của toàn bộ các ô quy hoạch thuộc loại đất này (m²).',
        'DT Trung Bình 1 Ô (m²)': 'Diện tích bình quân của 1 ô chức năng quy hoạch (m²).',
        'Diện Tích Ô (m²)': 'Diện tích của phần thửa đất nằm trong ô quy hoạch này (m²).',
        'Tỷ Lệ Chiếm Thửa (%)': 'Tỷ lệ % diện tích ô quy hoạch này so với tổng diện tích của toàn thửa đất.',
        'Tỷ Lệ Trên Tổng Khảo Sát (%)': 'Tỷ lệ % diện tích của loại đất này so với tổng diện tích toàn bộ dự án khảo sát.',
        'Có Nhà Ở': 'Đánh dấu thửa đất có chứa diện tích quy hoạch đất ở / dân cư (Có / Không).',
        'DT Nhà Ở (m²)': 'Diện tích phần đất thuộc quy hoạch đất ở / nhà ở trên thửa (m²).',
        'Tỷ Lệ Nhà Ở (%)': 'Tỷ lệ % diện tích đất ở / nhà ở so với tổng diện tích thửa đất.',
        'Thửa Có Nhà Ở': 'Số lượng thửa đất có chứa diện tích quy hoạch đất ở / dân cư.',
        'Có Công Trình XD': 'Đánh dấu thửa đất có chứa diện tích công trình xây dựng / TMDV / Công cộng / Cơ quan (Có / Không).',
        'DT Công Trình (m²)': 'Diện tích phần đất thuộc quy hoạch công trình xây dựng trên thửa (m²).',
        'Tỷ Lệ Công Trình (%)': 'Tỷ lệ % diện tích công trình xây dựng so với tổng diện tích thửa đất.',
        'Thửa Có Công Trình': 'Số lượng thửa đất có chứa diện tích quy hoạch công trình xây dựng.',
        'Có Nhà Ở Hoặc CT': 'Số lượng thửa đất có nhà ở HOẶC có công trình xây dựng (loại trừ đất giao thông thuần / đất khác).',
        'DT Lộ Giới (m²)': 'Diện tích phần thửa đất bị dính vào hành lang chỉ giới đường đỏ / quy hoạch mở đường (m²).',
        'DT Giao Thông Lộ Giới (m²)': 'Diện tích phần đất nằm trong lộ giới giao thông quy hoạch mở đường (m²).',
        'Tỷ Lệ Giao Thông (%)': 'Tỷ lệ % diện tích đất dính lộ giới giao thông so với diện tích thửa đất.',
        'Số Ô Xuất Hiện': 'Tổng số lượt ô phân khu chức năng thuộc loại đất này trên toàn bộ khảo sát.',
        'Số Thửa Tiếp Giáp': 'Số lượng thửa đất có tiếp giáp hoặc dính vào loại đất quy hoạch này.',
        'Số Thửa Khảo Sát': 'Tổng số lượng thửa đất được khảo sát tại Quận/Huyện này.',
        'Tỷ Lệ Thửa (%)': f'Tỷ lệ % số lượng thửa đất của Quận/Huyện so với toàn tuyến khảo sát {survey_id}.',
        'Thửa Có DA 1/500 & ĐCCB': 'Số lượng thửa đất có đồ án quy hoạch 1/500 hoặc điều chỉnh cục bộ tại địa bàn.',
        'Loại Quy Hoạch': 'Phân biệt Đồ án Quy hoạch chi tiết 1/500 (Dự án phát triển cụ thể) và Quyết định Điều chỉnh cục bộ (ĐCCB).',
        'Tên Dự Án / Quyết Định Điều Chỉnh': 'Tên đồ án quy hoạch 1/500 hoặc tên văn bản quyết định điều chỉnh quy hoạch.',
        'Tên Dự Án / Đồ Án Điều Chỉnh': 'Tên đồ án quy hoạch 1/500 hoặc tên văn bản quyết định điều chỉnh quy hoạch.',
        'Số Quyết Định': 'Số hiệu văn bản pháp lý phê duyệt quy hoạch do cơ quan có thẩm quyền ban hành.',
        'Số Quyết Định Phê Duyệt': 'Số hiệu văn bản pháp lý phê duyệt quy hoạch do cơ quan có thẩm quyền ban hành.',
        'Ngày Phê Duyệt': 'Ngày ban hành quyết định phê duyệt đồ án / điều chỉnh quy hoạch.',
        'Cơ Quan Phê Duyệt': 'Cơ quan nhà nước có thẩm quyền phê duyệt (UBND TP.HCM, Sở QHKT, UBND Quận...).',
        'Tỷ Lệ Thửa Thuộc Dự Án (%)': 'Tỷ lệ % diện tích thửa đất nằm trong ranh giới đồ án 1/500 hoặc quyết định điều chỉnh.'
    }

    def add_header_comment(cell, header_name):
        clean_name = str(header_name).strip()
        tooltip = HEADER_TOOLTIPS.get(clean_name, None)
        if tooltip:
            comment = Comment(f"📌 {clean_name}:\n{tooltip}", "Hệ Thống Quy Hoạch")
            comment.width = 280
            comment.height = 90
            cell.comment = comment

    # =========================================================================
    # SHEET 1: DASHBOARD & KPIS
    # =========================================================================
    print("1. Writing Sheet 1: Dashboard & KPIs...")
    ws1 = wb.create_sheet(title='📊 1. Dashboard & KPIs')
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells('A1:L1')
    ws1['A1'] = f"BÁO CÁO PHÂN TÍCH QUY HOẠCH ĐÔ THỊ - KHẢO SÁT CÔNG TRÌNH & NHÀ Ở {survey_id} (METRO SỐ 2)"
    ws1['A1'].font = font_title
    ws1['A1'].fill = fill_navy_hdr
    ws1['A1'].alignment = align_center
    ws1.row_dimensions[1].height = 30

    ws1.merge_cells('A2:L2')
    ws1['A2'] = f"Độ phân giải: {grid_m}m | Tổng số điểm quét: {total_points:,} | {len(df1):,} thửa đất | {len(df2):,} ô phân khu chức năng | Nguồn dữ liệu: SQHKT TP.HCM"
    ws1['A2'].font = font_subtitle
    ws1['A2'].fill = fill_navy_hdr
    ws1['A2'].alignment = align_center
    ws1.row_dimensions[2].height = 18

    kpis = [
        ('TỔNG SỐ THỬA ĐẤT', f"{len(df1):,}", '100% diện tích khảo sát', 'A', 'B'),
        ('TỔNG DIỆN TÍCH', f"{tot_survey_area/10000:.2f} ha", f"{tot_survey_area:,.1f} m²", 'C', 'D'),
        ('THỬA CÓ NHÀ Ở / DÂN CƯ', f"{tot_housing_parcels:,}", f"{tot_housing_parcels/len(df1)*100:.1f}% tổng số thửa", 'E', 'F'),
        ('THỬA CÓ CÔNG TRÌNH XD', f"{tot_facility_parcels:,}", f"{tot_facility_parcels/len(df1)*100:.1f}% TMDV / Cơ quan / CC", 'G', 'H'),
        ('CÓ NHÀ Ở HOẶC CÔNG TRÌNH', f"{tot_built_parcels:,}", f"{tot_built_parcels/len(df1)*100:.1f}% phục vụ khảo sát thực địa", 'I', 'J'),
        ('DÍNH LỘ GIỚI / GIAO THÔNG', f"{tot_traffic_parcels:,}", f"{tot_traffic_parcels/len(df1)*100:.1f}% tiếp giáp mở đường", 'K', 'L')
    ]

    for title, val, sub, col_start, col_end in kpis:
        ws1.merge_cells(f'{col_start}4:{col_end}4')
        ws1.merge_cells(f'{col_start}5:{col_end}5')
        ws1.merge_cells(f'{col_start}6:{col_end}6')
        
        ws1[f'{col_start}4'] = title
        ws1[f'{col_start}4'].font = font_kpi_label
        ws1[f'{col_start}4'].alignment = align_center
        ws1[f'{col_start}4'].fill = fill_kpi_bg
        
        ws1[f'{col_start}5'] = val
        ws1[f'{col_start}5'].font = font_kpi_num
        ws1[f'{col_start}5'].alignment = align_center
        ws1[f'{col_start}5'].fill = fill_kpi_bg
        
        ws1[f'{col_start}6'] = sub
        ws1[f'{col_start}6'].font = font_kpi_sub
        ws1[f'{col_start}6'].alignment = align_center
        ws1[f'{col_start}6'].fill = fill_kpi_bg
        
        for r in range(4, 7):
            for c_char in [col_start, col_end]:
                ws1[f'{c_char}{r}'].border = border_cell

    ws1.row_dimensions[4].height = 18
    ws1.row_dimensions[5].height = 26
    ws1.row_dimensions[6].height = 16

    num_districts = df1_enriched['Quận / Huyện'].nunique()
    ws1['A8'] = f"1. BẢNG PHÂN BỔ NHÀ Ở & CÔNG TRÌNH XÂY DỰNG THEO {num_districts} QUẬN / HUYỆN"
    ws1['A8'].font = font_section

    dist_agg = df1_enriched.groupby('Quận / Huyện').agg(
        Tong_Thua=('STT', 'count'),
        Tong_DT_m2=('Diện Tích Thửa (m²)', 'sum'),
        Thua_Nha_O=('Có Nhà Ở', lambda x: (x == 'Có').sum()),
        Thua_Cong_Trinh=('Có Công Trình XD', lambda x: (x == 'Có').sum()),
        Thua_Co_Nha_Hoac_CT=('Phân Loại Thực Địa', lambda x: (~x.str.contains('Giao Thông|Khác')).sum()),
        Thua_Lo_Gioi=('DT Giao Thông Lộ Giới (m²)', lambda x: (x > 0).sum()),
        Thua_1500=('Dự Án 1/500 & ĐCCB', lambda x: (x != 'Không').sum())
    ).reset_index()

    dist_agg['Ty_Le_Thua_%'] = dist_agg['Tong_Thua'] / len(df1) * 100
    dist_agg['Ty_Le_Nha_O_%'] = dist_agg['Thua_Nha_O'] / dist_agg['Tong_Thua'] * 100
    dist_agg['Ty_Le_CT_%'] = dist_agg['Thua_Cong_Trinh'] / dist_agg['Tong_Thua'] * 100
    dist_agg = dist_agg.sort_values(by='Tong_Thua', ascending=False)

    tbl1_headers = ['Quận / Huyện', 'Số Thửa Khảo Sát', 'Tỷ Lệ Thửa (%)', 'Tổng Diện Tích (m²)', 'Tổng DT (ha)', 'Thửa Có Nhà Ở', 'Tỷ Lệ Nhà Ở (%)', 'Thửa Có Công Trình', 'Tỷ Lệ Công Trình (%)', 'Có Nhà Ở Hoặc CT', 'Thửa Có DA 1/500 & ĐCCB']
    ws1.row_dimensions[9].height = 24
    for col_idx, h in enumerate(tbl1_headers, start=1):
        cell = ws1.cell(row=9, column=col_idx, value=h)
        cell.font = font_table_hdr
        cell.fill = fill_blue_hdr
        cell.alignment = align_center
        cell.border = border_header
        add_header_comment(cell, h)

    curr_row = 10
    for _, r in dist_agg.iterrows():
        ws1.row_dimensions[curr_row].height = 20
        ws1.cell(row=curr_row, column=1, value=r['Quận / Huyện']).alignment = align_left
        ws1.cell(row=curr_row, column=2, value=int(r['Tong_Thua'])).number_format = '#,##0'
        ws1.cell(row=curr_row, column=3, value=r['Ty_Le_Thua_%'] / 100).number_format = '0.00%'
        ws1.cell(row=curr_row, column=4, value=r['Tong_DT_m2']).number_format = '#,##0.0'
        ws1.cell(row=curr_row, column=5, value=r['Tong_DT_m2'] / 10000).number_format = '#,##0.00'
        ws1.cell(row=curr_row, column=6, value=int(r['Thua_Nha_O'])).number_format = '#,##0'
        ws1.cell(row=curr_row, column=7, value=r['Ty_Le_Nha_O_%'] / 100).number_format = '0.00%'
        ws1.cell(row=curr_row, column=8, value=int(r['Thua_Cong_Trinh'])).number_format = '#,##0'
        ws1.cell(row=curr_row, column=9, value=r['Ty_Le_CT_%'] / 100).number_format = '0.00%'
        ws1.cell(row=curr_row, column=10, value=int(r['Thua_Co_Nha_Hoac_CT'])).number_format = '#,##0'
        ws1.cell(row=curr_row, column=11, value=int(r['Thua_1500'])).number_format = '#,##0'
        
        for c in range(1, 12):
            cell = ws1.cell(row=curr_row, column=c)
            cell.font = font_cell
            cell.border = border_cell
            if c > 1:
                cell.alignment = align_right
            if curr_row % 2 == 1:
                cell.fill = fill_zebra
        curr_row += 1

    ws1.row_dimensions[curr_row].height = 22
    ws1.cell(row=curr_row, column=1, value="TỔNG CỘNG").alignment = align_left
    ws1.cell(row=curr_row, column=2, value=f"=SUM(B10:B{curr_row-1})").number_format = '#,##0'
    ws1.cell(row=curr_row, column=3, value=1.0).number_format = '0.00%'
    ws1.cell(row=curr_row, column=4, value=f"=SUM(D10:D{curr_row-1})").number_format = '#,##0.0'
    ws1.cell(row=curr_row, column=5, value=f"=SUM(E10:E{curr_row-1})").number_format = '#,##0.00'
    ws1.cell(row=curr_row, column=6, value=f"=SUM(F10:F{curr_row-1})").number_format = '#,##0'
    ws1.cell(row=curr_row, column=7, value=f"=F{curr_row}/B{curr_row}").number_format = '0.00%'
    ws1.cell(row=curr_row, column=8, value=f"=SUM(H10:H{curr_row-1})").number_format = '#,##0'
    ws1.cell(row=curr_row, column=9, value=f"=H{curr_row}/B{curr_row}").number_format = '0.00%'
    ws1.cell(row=curr_row, column=10, value=f"=SUM(J10:J{curr_row-1})").number_format = '#,##0'
    ws1.cell(row=curr_row, column=11, value=f"=SUM(K10:K{curr_row-1})").number_format = '#,##0'

    for c in range(1, 12):
        cell = ws1.cell(row=curr_row, column=c)
        cell.font = font_cell_bold
        cell.fill = fill_total_row
        cell.border = border_total

    # =========================================================================
    # SHEET 2: THỐNG KÊ TOÀN BỘ LOẠI ĐẤT THÔ
    # =========================================================================
    print(f"2. Writing Sheet 2: Thống kê {len(raw_types_df)} loại đất thô...")
    ws2 = wb.create_sheet(title=f'📚 2. Thống Kê {len(raw_types_df)} Loại Đất Thô')
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells('A1:J1')
    ws2['A1'] = f"BẢNG THỐNG KÊ TOÀN BỘ {len(raw_types_df)} LOẠI MỤC ĐÍCH SỬ DỤNG ĐẤT THÔ (KHÔNG GỘP - NGUYÊN BẢN SQHKT)"
    ws2['A1'].font = font_title
    ws2['A1'].fill = fill_navy_hdr
    ws2['A1'].alignment = align_center
    ws2.row_dimensions[1].height = 28

    ws2.merge_cells('A2:J2')
    ws2['A2'] = "Dữ liệu khảo sát thô ban đầu để xem chi tiết từng loại mục đích đất, số ô, số thửa, tổng diện tích m² và chỉ tiêu quy hoạch"
    ws2['A2'].font = font_subtitle
    ws2['A2'].fill = fill_navy_hdr
    ws2['A2'].alignment = align_center
    ws2.row_dimensions[2].height = 18

    s2_headers = [
        'STT', 'Tên Loại Đất Thô (SQHKT)', 'Nhóm Tham Khảo', 'Số Ô Xuất Hiện',
        'Số Thửa Tiếp Giáp', 'Tổng Diện Tích Ô (m²)', 'Tỷ Lệ Trên Tổng Khảo Sát (%)',
        'DT Trung Bình 1 Ô (m²)', 'Mật Độ XD Max (%)', 'HSSDĐ Max'
    ]
    ws2.row_dimensions[3].height = 24
    for c_idx, h in enumerate(s2_headers, start=1):
        cell = ws2.cell(row=3, column=c_idx, value=h)
        cell.font = font_table_hdr
        cell.fill = fill_teal_hdr
        cell.alignment = align_center
        cell.border = border_header
        add_header_comment(cell, h)

    for idx, r in raw_types_df.iterrows():
        row_num = idx + 4
        ws2.row_dimensions[row_num].height = 19
        ws2.cell(row=row_num, column=1, value=r['STT']).alignment = align_center
        ws2.cell(row=row_num, column=2, value=r['Tên Loại Đất Thô (SQHKT)']).alignment = align_left
        ws2.cell(row=row_num, column=3, value=r['Nhóm_Tham_Khảo']).alignment = align_left
        ws2.cell(row=row_num, column=4, value=int(r['Số_Ô'])).number_format = '#,##0'
        ws2.cell(row=row_num, column=5, value=int(r['Số_Thửa_Tiếp_Giáp'])).number_format = '#,##0'
        ws2.cell(row=row_num, column=6, value=r['Tổng_DT_Ô_m2']).number_format = '#,##0.0'
        ws2.cell(row=row_num, column=7, value=r['Tỷ Lệ DT Trên Tổng Khảo Sát (%)'] / 100).number_format = '0.00%'
        ws2.cell(row=row_num, column=8, value=r['DT_Trung_Bình_Ô_m2']).number_format = '#,##0.0'
        ws2.cell(row=row_num, column=9, value=r['Mật_Độ_Max'] if pd.notna(r['Mật_Độ_Max']) else '-').alignment = align_right
        ws2.cell(row=row_num, column=10, value=r['HSSDĐ_Max'] if pd.notna(r['HSSDĐ_Max']) else '-').alignment = align_right
        
        for c in range(1, 11):
            cell = ws2.cell(row=row_num, column=c)
            cell.font = font_cell
            cell.border = border_cell
            if 4 <= c <= 8:
                cell.alignment = align_right
            if row_num % 2 == 1:
                cell.fill = fill_zebra

    ws2.auto_filter.ref = f"A3:J{len(raw_types_df)+3}"
    ws2.freeze_panes = 'C4'

    # =========================================================================
    # SHEET 3: CHI TIẾT Ô QUY HOẠCH
    # =========================================================================
    print(f"3. Writing Sheet 3: Chi tiết {len(df2):,} ô quy hoạch...")
    ws3 = wb.create_sheet(title=f'🔍 3. Chi Tiết {len(df2):,} Ô QH')
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells('A1:O1')
    ws3['A1'] = f"DANH SÁCH TOÀN BỘ {len(df2):,} Ô PHÂN KHU CHỨC NĂNG QUY HOẠCH CHI TIẾT (ĐẦY ĐỦ NHÃN, MÃ Ô & DIỆN TÍCH)"
    ws3['A1'].font = font_title
    ws3['A1'].fill = fill_navy_hdr
    ws3['A1'].alignment = align_center
    ws3.row_dimensions[1].height = 28

    s3_headers = [
        'STT Ô', 'STT Thửa', 'Số Tờ', 'Số Thửa', 'Quận / Huyện', 'Phường / Xã',
        'Mã Ô Phố', 'Chức Năng Sử Dụng Đất (Thô)', 'Chức Năng Chi Tiết (SQHKT)',
        'Nhãn Ô Chi Tiết Đầy Đủ', 'Diện Tích Ô (m²)', 'Tỷ Lệ Chiếm Thửa (%)',
        'Tầng Cao Cho Phép', 'Mật Độ XD (%)', 'HSSDĐ'
    ]
    ws3.row_dimensions[2].height = 24
    for c_idx, h in enumerate(s3_headers, start=1):
        cell = ws3.cell(row=2, column=c_idx, value=h)
        cell.font = font_table_hdr
        cell.fill = fill_blue_hdr
        cell.alignment = align_center
        cell.border = border_header
        add_header_comment(cell, h)

    for idx, (_, r) in enumerate(df2.iterrows(), start=1):
        row_num = idx + 2
        ws3.row_dimensions[row_num].height = 18
        ws3.cell(row=row_num, column=1, value=idx).alignment = align_center
        ws3.cell(row=row_num, column=2, value=int(r['STT Thửa'])).alignment = align_center
        ws3.cell(row=row_num, column=3, value=r['Số Tờ']).alignment = align_center
        ws3.cell(row=row_num, column=4, value=r['Số Thửa']).alignment = align_center
        ws3.cell(row=row_num, column=5, value=r['Quận / Huyện']).alignment = align_left
        ws3.cell(row=row_num, column=6, value=r['Phường / Xã']).alignment = align_left
        ws3.cell(row=row_num, column=7, value=str(r['Mã Ô Phố'])).alignment = align_center
        ws3.cell(row=row_num, column=8, value=r['Chức Năng Sử Dụng Đất']).alignment = align_left
        ws3.cell(row=row_num, column=9, value=r['Chức Năng Chi Tiết (SQHKT)']).alignment = align_left
        ws3.cell(row=row_num, column=10, value=r['Nhãn Ô Chi Tiết']).alignment = align_left
        ws3.cell(row=row_num, column=11, value=r['Diện Tích Ô (m²)']).number_format = '#,##0.0'
        ws3.cell(row=row_num, column=12, value=r['Tỷ Lệ Chiếm Thửa (%)']).number_format = '0.0%'
        ws3.cell(row=row_num, column=13, value=r['Tầng Cao Cho Phép (tầng)'] if pd.notna(r['Tầng Cao Cho Phép (tầng)']) else '-').alignment = align_center
        ws3.cell(row=row_num, column=14, value=r['Mật Độ XD Tối Đa (%)'] if pd.notna(r['Mật Độ XD Tối Đa (%)']) else '-').alignment = align_right
        ws3.cell(row=row_num, column=15, value=r['Hệ Số Sử Dụng Đất (HSSDĐ)'] if pd.notna(r['Hệ Số Sử Dụng Đất (HSSDĐ)']) else '-').alignment = align_right
        
        for c in range(1, 16):
            cell = ws3.cell(row=row_num, column=c)
            cell.font = font_cell
            cell.border = border_cell
            if c in [11, 12]:
                cell.alignment = align_right
            if row_num % 2 == 1:
                cell.fill = fill_zebra

    ws3.auto_filter.ref = f"A2:O{len(df2)+2}"
    ws3.freeze_panes = 'E3'

    # =========================================================================
    # SHEET 4: THỬA CÓ ĐẤT Ở - NHÀ Ở
    # =========================================================================
    df_housing_parcels = df1_enriched[df1_enriched['Có Nhà Ở'] == 'Có'].copy()
    print(f"4. Writing Sheet 4: Thửa có đất ở - nhà ở ({len(df_housing_parcels):,} thửa)...")
    ws4 = wb.create_sheet(title='🏡 4. Thửa Có Đất Ở - Nhà Ở')
    ws4.views.sheetView[0].showGridLines = True

    ws4.merge_cells('A1:O1')
    ws4['A1'] = f"DANH SÁCH {len(df_housing_parcels):,} THỬA ĐẤT CÓ CHỨC NĂNG ĐẤT Ở / NHÀ Ở / DÂN CƯ HIỆN HỮU"
    ws4['A1'].font = font_title
    ws4['A1'].fill = fill_navy_hdr
    ws4['A1'].alignment = align_center
    ws4.row_dimensions[1].height = 28

    s4_headers = [
        'STT', 'Mã Thửa Đất', 'Số Tờ', 'Số Thửa', 'Quận / Huyện', 'Phường / Xã',
        'Tổng DT Thửa (m²)', 'Loại Đất Nhà Ở Thô', 'DT Nhà Ở (m²)', 'Tỷ Lệ Nhà Ở (%)',
        'Tầng Cao Max', 'Mật Độ XD Max (%)', 'HSSDĐ Max', 'Dự Án 1/500 & ĐCCB', 'Lộ Giới Tiếp Giáp'
    ]
    ws4.row_dimensions[2].height = 24
    for c_idx, h in enumerate(s4_headers, start=1):
        cell = ws4.cell(row=2, column=c_idx, value=h)
        cell.font = font_table_hdr
        cell.fill = fill_blue_hdr
        cell.alignment = align_center
        cell.border = border_header
        add_header_comment(cell, h)

    for idx, (_, r) in enumerate(df_housing_parcels.iterrows(), start=1):
        row_num = idx + 2
        ws4.row_dimensions[row_num].height = 18
        ws4.cell(row=row_num, column=1, value=idx).alignment = align_center
        ws4.cell(row=row_num, column=2, value=str(r['Mã Thửa Đất'])).alignment = align_left
        ws4.cell(row=row_num, column=3, value=r['Số Tờ']).alignment = align_center
        ws4.cell(row=row_num, column=4, value=r['Số Thửa']).alignment = align_center
        ws4.cell(row=row_num, column=5, value=r['Quận / Huyện']).alignment = align_left
        ws4.cell(row=row_num, column=6, value=r['Phường / Xã']).alignment = align_left
        ws4.cell(row=row_num, column=7, value=r['Diện Tích Thửa (m²)']).number_format = '#,##0.0'
        ws4.cell(row=row_num, column=8, value=r['Loại Đất Nhà Ở Thô']).alignment = align_left
        ws4.cell(row=row_num, column=9, value=r['DT Nhà Ở (m²)']).number_format = '#,##0.0'
        ws4.cell(row=row_num, column=10, value=r['Tỷ Lệ Nhà Ở (%)'] / 100).number_format = '0.0%'
        ws4.cell(row=row_num, column=11, value=r['Tầng Cao Nhà Ở'] if pd.notna(r['Tầng Cao Nhà Ở']) else '-').alignment = align_center
        ws4.cell(row=row_num, column=12, value=r['Mật Độ Nhà Ở (%)'] if pd.notna(r['Mật Độ Nhà Ở (%)']) else '-').alignment = align_right
        ws4.cell(row=row_num, column=13, value=r['HSSDĐ Nhà Ở'] if pd.notna(r['HSSDĐ Nhà Ở']) else '-').alignment = align_right
        ws4.cell(row=row_num, column=14, value=r['Dự Án 1/500 & ĐCCB']).alignment = align_center
        ws4.cell(row=row_num, column=15, value=r['Lộ Giới Tuyến Đường Tiếp Giáp']).alignment = align_left
        
        for c in range(1, 16):
            cell = ws4.cell(row=row_num, column=c)
            cell.font = font_cell
            cell.border = border_cell
            if c in [7, 9, 10]:
                cell.alignment = align_right
            if row_num % 2 == 1:
                cell.fill = fill_zebra

    ws4.auto_filter.ref = f"A2:O{len(df_housing_parcels)+2}"
    ws4.freeze_panes = 'E3'

    # =========================================================================
    # SHEET 5: THỬA CÓ CÔNG TRÌNH XÂY DỰNG
    # =========================================================================
    df_facility_parcels = df1_enriched[df1_enriched['Có Công Trình XD'] == 'Có'].sort_values(by='DT Công Trình (m²)', ascending=False).copy()
    print(f"5. Writing Sheet 5: Thửa có công trình xây dựng ({len(df_facility_parcels):,} thửa)...")
    ws5 = wb.create_sheet(title='🏢 5. Thửa Có Công Trình XD')
    ws5.views.sheetView[0].showGridLines = True

    ws5.merge_cells('A1:P1')
    ws5['A1'] = f"DANH SÁCH {len(df_facility_parcels):,} THỬA ĐẤT CÓ CÔNG TRÌNH XÂY DỰNG / TMDV / CƠ QUAN / CÔNG CỘNG / PHỨC HỢP / Y TẾ / GIÁO DỤC"
    ws5['A1'].font = font_title
    ws5['A1'].fill = fill_navy_hdr
    ws5['A1'].alignment = align_center
    ws5.row_dimensions[1].height = 28

    s5_headers = [
        'STT', 'Mã Thửa Đất', 'Số Tờ', 'Số Thửa', 'Quận / Huyện', 'Phường / Xã',
        'Tổng DT Thửa (m²)', 'Phân Loại Công Trình', 'Tên Loại Đất Công Trình Thô',
        'DT Công Trình (m²)', 'Tỷ Lệ Công Trình (%)', 'Tầng Cao Max', 'Mật Độ XD Max (%)', 'HSSDĐ Max',
        'Dự Án 1/500 & ĐCCB', 'Lộ Giới Tiếp Giáp'
    ]
    ws5.row_dimensions[2].height = 24
    for c_idx, h in enumerate(s5_headers, start=1):
        cell = ws5.cell(row=2, column=c_idx, value=h)
        cell.font = font_table_hdr
        cell.fill = fill_purple_hdr
        cell.alignment = align_center
        cell.border = border_header
        add_header_comment(cell, h)

    for idx, (_, r) in enumerate(df_facility_parcels.iterrows(), start=1):
        row_num = idx + 2
        ws5.row_dimensions[row_num].height = 18
        ws5.cell(row=row_num, column=1, value=idx).alignment = align_center
        ws5.cell(row=row_num, column=2, value=str(r['Mã Thửa Đất'])).alignment = align_left
        ws5.cell(row=row_num, column=3, value=r['Số Tờ']).alignment = align_center
        ws5.cell(row=row_num, column=4, value=r['Số Thửa']).alignment = align_center
        ws5.cell(row=row_num, column=5, value=r['Quận / Huyện']).alignment = align_left
        ws5.cell(row=row_num, column=6, value=r['Phường / Xã']).alignment = align_left
        ws5.cell(row=row_num, column=7, value=r['Diện Tích Thửa (m²)']).number_format = '#,##0.0'
        ws5.cell(row=row_num, column=8, value=r['Phân Loại Công Trình']).alignment = align_left
        ws5.cell(row=row_num, column=9, value=r['Tên Đất Công Trình Thô']).alignment = align_left
        ws5.cell(row=row_num, column=10, value=r['DT Công Trình (m²)']).number_format = '#,##0.0'
        ws5.cell(row=row_num, column=11, value=r['Tỷ Lệ Công Trình (%)'] / 100).number_format = '0.0%'
        ws5.cell(row=row_num, column=12, value=r['Tầng Cao CT'] if pd.notna(r['Tầng Cao CT']) else '-').alignment = align_center
        ws5.cell(row=row_num, column=13, value=r['Mật Độ CT (%)'] if pd.notna(r['Mật Độ CT (%)']) else '-').alignment = align_right
        ws5.cell(row=row_num, column=14, value=r['HSSDĐ CT'] if pd.notna(r['HSSDĐ CT']) else '-').alignment = align_right
        ws5.cell(row=row_num, column=15, value=r['Dự Án 1/500 & ĐCCB']).alignment = align_center
        ws5.cell(row=row_num, column=16, value=r['Lộ Giới Tuyến Đường Tiếp Giáp']).alignment = align_left
        
        for c in range(1, 17):
            cell = ws5.cell(row=row_num, column=c)
            cell.font = font_cell
            cell.border = border_cell
            if c in [7, 10, 11]:
                cell.alignment = align_right
            if row_num % 2 == 1:
                cell.fill = fill_zebra
            if c == 8:
                cell.fill = fill_highlight_blue

    ws5.auto_filter.ref = f"A2:P{len(df_facility_parcels)+2}"
    ws5.freeze_panes = 'E3'

    # =========================================================================
    # SHEET 6: TOÀN BỘ THỬA CÓ NHÀ Ở HOẶC CÔNG TRÌNH
    # =========================================================================
    df_all_built = df1_enriched[(df1_enriched['Có Nhà Ở'] == 'Có') | (df1_enriched['Có Công Trình XD'] == 'Có')].copy()
    print(f"6. Writing Sheet 6: Toàn bộ thửa có nhà ở hoặc công trình ({len(df_all_built):,} thửa)...")
    ws6 = wb.create_sheet(title=f'🏗️ 6. Toàn Bộ Nhà Ở & CT ({len(df_all_built)})')
    ws6.views.sheetView[0].showGridLines = True

    ws6.merge_cells('A1:P1')
    ws6['A1'] = f"DANH SÁCH TỔNG HỢP {len(df_all_built):,} THỬA ĐẤT CÓ NHÀ Ở HOẶC CÔNG TRÌNH XÂY DỰNG (PHỤC VỤ KHẢO SÁT THỰC ĐỊA)"
    ws6['A1'].font = font_title
    ws6['A1'].fill = fill_navy_hdr
    ws6['A1'].alignment = align_center
    ws6.row_dimensions[1].height = 28

    s6_headers = [
        'STT Thửa', 'Số Tờ', 'Số Thửa', 'Mã Thửa Đất', 'Quận / Huyện', 'Phường / Xã',
        'Tổng DT Thửa (m²)', 'Phân Loại Thực Địa', 'DT Nhà Ở (m²)', 'Tỷ Lệ Nhà Ở (%)',
        'DT Công Trình (m²)', 'Tỷ Lệ Công Trình (%)', 'DT Lộ Giới (m²)', 'Dự Án 1/500 & ĐCCB',
        'Cơ Cấu Quy Hoạch Thô (SQHKT)', 'Toạ Độ Khảo Sát'
    ]
    ws6.row_dimensions[2].height = 24
    for c_idx, h in enumerate(s6_headers, start=1):
        cell = ws6.cell(row=2, column=c_idx, value=h)
        cell.font = font_table_hdr
        cell.fill = fill_teal_hdr
        cell.alignment = align_center
        cell.border = border_header
        add_header_comment(cell, h)

    for idx, (_, r) in enumerate(df_all_built.iterrows(), start=1):
        row_num = idx + 2
        ws6.row_dimensions[row_num].height = 18
        ws6.cell(row=row_num, column=1, value=int(r['STT'])).alignment = align_center
        ws6.cell(row=row_num, column=2, value=r['Số Tờ']).alignment = align_center
        ws6.cell(row=row_num, column=3, value=r['Số Thửa']).alignment = align_center
        ws6.cell(row=row_num, column=4, value=str(r['Mã Thửa Đất'])).alignment = align_left
        ws6.cell(row=row_num, column=5, value=r['Quận / Huyện']).alignment = align_left
        ws6.cell(row=row_num, column=6, value=r['Phường / Xã']).alignment = align_left
        ws6.cell(row=row_num, column=7, value=r['Diện Tích Thửa (m²)']).number_format = '#,##0.0'
        ws6.cell(row=row_num, column=8, value=r['Phân Loại Thực Địa']).alignment = align_left
        ws6.cell(row=row_num, column=9, value=r['DT Nhà Ở (m²)']).number_format = '#,##0.0'
        ws6.cell(row=row_num, column=10, value=r['Tỷ Lệ Nhà Ở (%)'] / 100).number_format = '0.0%'
        ws6.cell(row=row_num, column=11, value=r['DT Công Trình (m²)']).number_format = '#,##0.0'
        ws6.cell(row=row_num, column=12, value=r['Tỷ Lệ Công Trình (%)'] / 100).number_format = '0.0%'
        ws6.cell(row=row_num, column=13, value=r['DT Giao Thông Lộ Giới (m²)']).number_format = '#,##0.0'
        ws6.cell(row=row_num, column=14, value=r['Dự Án 1/500 & ĐCCB']).alignment = align_center
        ws6.cell(row=row_num, column=15, value=r['Cơ Cấu Quy Hoạch Sử Dụng Đất']).alignment = align_left
        ws6.cell(row=row_num, column=16, value=r['Toạ Độ Khảo Sát (Kinh độ, Vĩ độ)']).alignment = align_left
        
        for c in range(1, 17):
            cell = ws6.cell(row=row_num, column=c)
            cell.font = font_cell
            cell.border = border_cell
            if c in [7, 9, 10, 11, 12, 13]:
                cell.alignment = align_right
            if row_num % 2 == 1:
                cell.fill = fill_zebra

    ws6.auto_filter.ref = f"A2:P{len(df_all_built)+2}"
    ws6.freeze_panes = 'E3'

    # =========================================================================
    # SHEET 7: DỰ ÁN 1-500 & ĐCCB
    # =========================================================================
    print(f"7. Writing Sheet 7: Dự án 1/500 & điều chỉnh cục bộ ({len(df4):,} bản ghi)...")
    ws7 = wb.create_sheet(title='🏛️ 7. Dự Án 1-500 & DCCB')
    ws7.views.sheetView[0].showGridLines = True

    ws7.merge_cells('A1:L1')
    ws7['A1'] = f"DANH MỤC {len(df4):,} BẢN GHI THUỘC ĐỒ ÁN QUY HOẠCH 1/500 & ĐIỀU CHỈNH CỤC BỘ"
    ws7['A1'].font = font_title
    ws7['A1'].fill = fill_navy_hdr
    ws7['A1'].alignment = align_center
    ws7.row_dimensions[1].height = 28

    s7_headers = [
        'STT', 'Mã Thửa Đất', 'Số Tờ', 'Số Thửa', 'Quận / Huyện', 'Phường / Xã',
        'Loại Quy Hoạch', 'Tên Dự Án / Quyết Định Điều Chỉnh', 'Số Quyết Định',
        'Ngày Phê Duyệt', 'Cơ Quan Phê Duyệt', 'Tỷ Lệ Thửa Thuộc Dự Án (%)'
    ]
    ws7.row_dimensions[2].height = 24
    for c_idx, h in enumerate(s7_headers, start=1):
        cell = ws7.cell(row=2, column=c_idx, value=h)
        cell.font = font_table_hdr
        cell.fill = fill_purple_hdr
        cell.alignment = align_center
        cell.border = border_header
        add_header_comment(cell, h)

    for idx, (_, r) in enumerate(df4.iterrows(), start=1):
        row_num = idx + 2
        ws7.row_dimensions[row_num].height = 18
        ws7.cell(row=row_num, column=1, value=idx).alignment = align_center
        ws7.cell(row=row_num, column=2, value=str(r['Mã Thửa Đất'])).alignment = align_left
        ws7.cell(row=row_num, column=3, value=r['Số Tờ']).alignment = align_center
        ws7.cell(row=row_num, column=4, value=r['Số Thửa']).alignment = align_center
        ws7.cell(row=row_num, column=5, value=r['Quận / Huyện']).alignment = align_left
        ws7.cell(row=row_num, column=6, value=r['Phường / Xã']).alignment = align_left
        ws7.cell(row=row_num, column=7, value=r['Loại Quy Hoạch']).alignment = align_center
        ws7.cell(row=row_num, column=8, value=r['Tên Dự Án / Đồ Án Điều Chỉnh']).alignment = align_left
        ws7.cell(row=row_num, column=9, value=r['Số Quyết Định Phê Duyệt'] if pd.notna(r['Số Quyết Định Phê Duyệt']) else '-').alignment = align_left
        ws7.cell(row=row_num, column=10, value=str(r['Ngày Phê Duyệt']) if pd.notna(r['Ngày Phê Duyệt']) else '-').alignment = align_center
        ws7.cell(row=row_num, column=11, value=r['Cơ Quan Phê Duyệt'] if pd.notna(r['Cơ Quan Phê Duyệt']) else '-').alignment = align_left
        
        pct_val = r['Tỷ Lệ Thửa Thuộc Dự Án (%)']
        if pd.notna(pct_val):
            ws7.cell(row=row_num, column=12, value=float(pct_val) if float(pct_val) <= 1.0 else float(pct_val)/100).number_format = '0.0%'
        else:
            ws7.cell(row=row_num, column=12, value='-').alignment = align_center
        
        for c in range(1, 13):
            cell = ws7.cell(row=row_num, column=c)
            cell.font = font_cell
            cell.border = border_cell
            if c == 12:
                cell.alignment = align_right
            if row_num % 2 == 1:
                cell.fill = fill_zebra
            if '1/500' in str(r['Loại Quy Hoạch']):
                if c == 7:
                    cell.fill = fill_highlight_amber

    ws7.auto_filter.ref = f"A2:L{len(df4)+2}"
    ws7.freeze_panes = 'E3'

    # Auto column widths
    print("8. Auto-adjusting column widths...")
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row in [1, 2] and col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P']:
                    continue
                if cell.value:
                    val_str = str(cell.value)
                    max_len = max(max_len, len(val_str))
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 65)

    ws1.column_dimensions['A'].width = 30
    ws1.column_dimensions['B'].width = 18
    ws1.column_dimensions['C'].width = 18
    ws1.column_dimensions['D'].width = 22
    ws1.column_dimensions['E'].width = 16
    ws1.column_dimensions['F'].width = 18
    ws1.column_dimensions['G'].width = 38

    print(f"9. Saving workbook to: {output_excel}")
    wb.save(output_excel)
    print(f"SUCCESS: Report for {survey_id} generated at {output_excel}!")

if __name__ == '__main__':
    survey = sys.argv[1] if len(sys.argv) > 1 else 'KS003'
    generate_report(survey)
