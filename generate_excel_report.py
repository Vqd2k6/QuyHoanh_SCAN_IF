import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import numpy as np
import unicodedata

excel_path = 'data/output/KS002/KS002.xlsx'
output_excel = 'data/output/KS002/Bao_Cao_Phan_Tich_Quy_Hoach_KS002.xlsx'

print("1. Loading raw sheets from:", excel_path)
df1 = pd.read_excel(excel_path, sheet_name=0, header=3)
df2 = pd.read_excel(excel_path, sheet_name=1, header=3)
df3 = pd.read_excel(excel_path, sheet_name=2, header=3)
df4 = pd.read_excel(excel_path, sheet_name=3, header=3)

def normalize(s):
    if not isinstance(s, str):
        return ''
    return unicodedata.normalize('NFC', s).strip()

df2['norm_name'] = df2['Chức Năng Sử Dụng Đất'].apply(normalize)

def get_category_info(name):
    n = name.lower()
    # Green spaces
    if any(k in n for k in ['cây xanh', 'công viên', 'cv -', 'cv-', 'thể dục thể thao']):
        return ('5. Cây Xanh, Công Viên & TDTT', 'Đất công viên cây xanh, vườn hoa, thể thao, cây xanh cách ly', 'Hạn chế / Không cấp phép XD mới, nguy cơ thu hồi', '🔴 RỦI RO CAO (Dính Cây Xanh)')
    # Residential
    if any(k in n for k in ['đất ở', 'dân cư', 'nhà ở', 'nhóm nhà ở']):
        if 'không cho xây dựng mới' in n:
            return ('1. Đất Ở / Thổ Cư & Dân Cư', 'Đất ở hiện hữu giữ nguyên hiện trạng', 'Giữ nguyên hiện trạng, không cấp phép XD mới', '🟡 CẦN THẨM TRA (Hạn Chế XD)')
        elif 'xây dựng mới' in n or 'dự kiến' in n:
            return ('1. Đất Ở / Thổ Cư & Dân Cư', 'Đất ở quy hoạch xây dựng mới / dự kiến', 'Cấp phép xây dựng theo đồ án quy hoạch chi tiết', '🟢 AN TOÀN (Quy Hoạch Mới)')
        else:
            return ('1. Đất Ở / Thổ Cư & Dân Cư', 'Đất ở đô thị / thổ cư (hiện hữu, cải tạo, chỉnh trang)', 'Được cấp phép xây dựng nhà ở đô thị chính thức', '🟢 AN TOÀN (Thổ Cư Chuẩn)')
    # Mixed use
    if any(k in n for k in ['phức hợp', 'hỗn hợp']):
        return ('2. Đất Phức Hợp / Hỗn Hợp', 'Đất sử dụng đa chức năng (Ở kết hợp Thương mại, Dịch vụ, Văn phòng)', 'Được cấp phép theo đồ án tỷ lệ cơ cấu chức năng dự án', '🔵 ĐẤT PHỨC HỢP / HỖN HỢP')
    # Commercial & Office
    if any(k in n for k in ['thương mại', 'tmdv', 'sxkd', 'cơ quan']):
        return ('3. Thương Mại - Dịch Vụ & Cơ Quan', 'Đất công trình thương mại, dịch vụ, trụ sở cơ quan, SXKD', 'Cấp phép công trình thương mại, văn phòng, dịch vụ thương mại', '🔵 THƯƠNG MẠI - DỊCH VỤ')
    # Public & Institutional
    if any(k in n for k in ['công cộng', 'giáo dục', 'trường học', 'y tế', 'hành chính', 'hành chánh', 'tôn giáo', 'di tích', 'quảng trường']):
        return ('4. Công Trình Công Cộng, Y Tế, GD & Tôn Giáo', 'Đất công trình an sinh, bệnh viện, trường học, hành chính, tôn giáo', 'Dành riêng cho công trình công cộng, không được XD nhà ở riêng lẻ', '🔴 RỦI RO CAO (Dính Công Cộng)')
    # Transport & Infrastructure & Water
    if any(k in n for k in ['giao thông', 'ga depot', 'depot', 'mặt nước', 'kênh', 'rạch', 'sông', 'hạ tầng', 'cấp nước', 'kho tàng']):
        return ('6. Giao Thông, Hạ Tầng & Mặt Nước', 'Đất lộ giới đường bộ, ga depot metro, hạ tầng kỹ thuật, hành lang sông rạch', 'Nằm trong lộ giới mở đường / hành lang kỹ thuật, bị trừ lộ giới', '⚠️ LỘ GIỚI / HẠ TẦNG')
    # Industry & Defense & Others
    if any(k in n for k in ['công nghiệp', 'tiểu thủ', 'quân sự', 'quốc phòng']):
        return ('7. Công Nghiệp & Quốc Phòng', 'Đất an ninh quốc phòng, sản xuất công nghiệp, kho bãi', 'Đất chuyên dùng Nhà nước / Sản xuất công nghiệp', '⚠️ ĐẤT CHUYÊN DÙNG')
    return ('8. Đất Khác', 'Đất khác theo phân loại quy hoạch', 'Cần kiểm tra hồ sơ địa chính chi tiết', 'ℹ️ CẦN THẨM TRA')

cat_res = df2['norm_name'].apply(get_category_info)
df2['Nhom_Quy_Hoach'] = [r[0] for r in cat_res]
df2['Y_Nghia_Su_Dung'] = [r[1] for r in cat_res]
df2['Kha_Nang_Xay_Dung'] = [r[2] for r in cat_res]
df2['Muc_Do_Rui_Ro'] = [r[3] for r in cat_res]

# Group parcels
parcel_cats = df2.groupby('STT Thửa')['Nhom_Quy_Hoach'].unique().to_dict()
parcel_green_details = df2[df2['Nhom_Quy_Hoach'] == '5. Cây Xanh, Công Viên & TDTT'].groupby('STT Thửa').agg(
    Loai_Cay_Xanh=('norm_name', lambda x: ', '.join(x.unique())),
    Ty_Le_Cay_Xanh=('Tỷ Lệ Chiếm Thửa (%)', 'sum'),
    DT_Cay_Xanh=('Diện Tích Ô (m²)', 'sum')
).reset_index()
green_dict = {row['STT Thửa']: row for _, row in parcel_green_details.iterrows()}

parcel_res_details = df2[df2['Nhom_Quy_Hoach'] == '1. Đất Ở / Thổ Cư & Dân Cư'].groupby('STT Thửa').agg(
    Loai_Dat_O=('norm_name', lambda x: ', '.join(x.unique())),
    Ty_Le_Dat_O=('Tỷ Lệ Chiếm Thửa (%)', 'sum'),
    DT_Dat_O=('Diện Tích Ô (m²)', 'sum')
).reset_index()
res_dict = {row['STT Thửa']: row for _, row in parcel_res_details.iterrows()}

# Parcels with 1/500 from df4
parcels_with_1500 = set(df4[df4['Loại Quy Hoạch'] == 'Quy hoạch chi tiết 1/500']['STT Thửa'].unique())
parcels_with_dccb = set(df4[df4['Loại Quy Hoạch'] == 'Điều chỉnh cục bộ (DCCB)']['STT Thửa'].unique())

def enrich_parcel(row):
    stt = row['STT']
    cats = parcel_cats.get(stt, [])
    has_res = any('1. Đất Ở' in c for c in cats)
    has_mixed = any('2. Đất Phức Hợp' in c for c in cats)
    has_green = any('5. Cây Xanh' in c for c in cats)
    has_public = any('4. Công Trình' in c for c in cats)
    has_traffic = any('6. Giao Thông' in c for c in cats)
    
    has_1500_exact = stt in parcels_with_1500
    has_dccb_exact = stt in parcels_with_dccb
    
    if has_green:
        badge = '🔴 Dính Cây Xanh'
        status_safe = 'Rủi ro Cây Xanh'
    elif has_public:
        badge = '🔴 Dính Công Cộng'
        status_safe = 'Rủi ro Công Cộng'
    elif has_res and not has_green and not has_public:
        badge = '🟢 Thổ Cư An Toàn'
        status_safe = 'Thổ Cư Sạch'
    elif has_mixed:
        badge = '🔵 Phức Hợp / Hỗn Hợp'
        status_safe = 'Phức Hợp'
    else:
        badge = '⚪ Khác / Chuyên Dùng'
        status_safe = 'Khác'
        
    res_info = res_dict.get(stt, None)
    green_info = green_dict.get(stt, None)
    
    ty_le_o = round(res_info['Ty_Le_Dat_O'] * 100, 2) if res_info is not None else 0.0
    dt_o = round(res_info['DT_Dat_O'], 2) if res_info is not None else 0.0
    
    ty_le_cx = round(green_info['Ty_Le_Cay_Xanh'] * 100, 2) if green_info is not None else 0.0
    dt_cx = round(green_info['DT_Cay_Xanh'], 2) if green_info is not None else 0.0
    loai_cx = green_info['Loai_Cay_Xanh'] if green_info is not None else 'Không'

    status_1500 = 'Có QH 1/500' if has_1500_exact else ('Có ĐCCB' if has_dccb_exact else 'Không')

    return pd.Series({
        'Phân Loại Quy Hoạch': badge,
        'Trạng Thái An Toàn': status_safe,
        'Có Đất Ở': 'Có' if has_res else 'Không',
        'Tỷ Lệ Đất Ở (%)': ty_le_o,
        'Diện Tích Đất Ở (m²)': dt_o,
        'Có Cây Xanh': 'CÓ (DÍNH)' if has_green else 'Không',
        'Tỷ Lệ Cây Xanh (%)': ty_le_cx,
        'Diện Tích Cây Xanh (m²)': dt_cx,
        'Loại QH Cây Xanh Dính Phải': loai_cx,
        'Dính Lộ Giới / Giao Thông': 'Có' if has_traffic else 'Không',
        'Dự Án 1/500 & ĐCCB': status_1500
    })

enriched = df1.apply(enrich_parcel, axis=1)
df1_enriched = pd.concat([df1, enriched], axis=1)

# Build Catalog DataFrame (72 types)
cat_df = df2.groupby('norm_name').agg(
    Nhom_Quy_Hoach=('Nhom_Quy_Hoach', 'first'),
    Y_Nghia=('Y_Nghia_Su_Dung', 'first'),
    Kha_Nang_XD=('Kha_Nang_Xay_Dung', 'first'),
    Muc_Do_Rui_Ro=('Muc_Do_Rui_Ro', 'first'),
    So_Lan_Xuat_Hien=('STT Thửa', 'count'),
    So_Thua_Duy_Nhat=('STT Thửa', 'nunique'),
    Tong_DT_O_m2=('Diện Tích Ô (m²)', 'sum'),
    Tang_Cao_Max=('Tầng Cao Cho Phép (tầng)', lambda x: pd.to_numeric(x, errors='coerce').max()),
    Mat_Do_Max=('Mật Độ XD Tối Đa (%)', lambda x: pd.to_numeric(x, errors='coerce').max()),
    HSSDD_Max=('Hệ Số Sử Dụng Đất (HSSDĐ)', lambda x: pd.to_numeric(x, errors='coerce').max())
).reset_index().rename(columns={'norm_name': 'Tên Chức Năng SQHKT'})

cat_df = cat_df.sort_values(by=['Nhom_Quy_Hoach', 'So_Thua_Duy_Nhat'], ascending=[True, False]).reset_index(drop=True)
cat_df.insert(0, 'STT', range(1, len(cat_df) + 1))

# Excel Building with openpyxl
print("2. Constructing styled Excel workbook...")
wb = openpyxl.Workbook()
wb.remove(wb.active)

FONT_NAME = 'Arial'
font_title = Font(name=FONT_NAME, size=15, bold=True, color='FFFFFF')
font_subtitle = Font(name=FONT_NAME, size=10, italic=True, color='E2E8F0')
font_section = Font(name=FONT_NAME, size=11, bold=True, color='1E3A8A')
font_table_hdr = Font(name=FONT_NAME, size=10, bold=True, color='FFFFFF')
font_kpi_num = Font(name=FONT_NAME, size=16, bold=True, color='1E3A8A')
font_kpi_label = Font(name=FONT_NAME, size=9, bold=True, color='475569')
font_kpi_sub = Font(name=FONT_NAME, size=8, italic=True, color='64748B')
font_cell = Font(name=FONT_NAME, size=9)
font_cell_bold = Font(name=FONT_NAME, size=9, bold=True)

fill_navy_hdr = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
fill_blue_hdr = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
fill_teal_hdr = PatternFill(start_color='0D9488', end_color='0D9488', fill_type='solid')
fill_purple_hdr = PatternFill(start_color='4338CA', end_color='4338CA', fill_type='solid')
fill_kpi_bg = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
fill_highlight_green = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')
fill_highlight_red = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
fill_highlight_amber = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
fill_total_row = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')

thin_side = Side(border_style='thin', color='CBD5E1')
thick_bottom = Side(border_style='medium', color='1E3A8A')
double_bottom = Side(border_style='double', color='1E3A8A')

border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
border_header = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thick_bottom)
border_total = Border(left=thin_side, right=thin_side, top=thin_side, bottom=double_bottom)

align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
align_right = Alignment(horizontal='right', vertical='center')

# ==========================================
# SHEET 1: DASHBOARD & TỔNG QUAN
# ==========================================
ws1 = wb.create_sheet(title='📊 1. Dashboard & Tổng Quan')
ws1.views.sheetView[0].showGridLines = True

ws1.merge_cells('A1:L1')
ws1['A1'] = "BÁO CÁO PHÂN TÍCH TỔNG HỢP QUY HOẠCH ĐÔ THỊ TP. HỒ CHÍ MINH - KHẢO SÁT KS002"
ws1['A1'].font = font_title
ws1['A1'].fill = fill_navy_hdr
ws1['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws1.row_dimensions[1].height = 32

ws1.merge_cells('A2:L2')
ws1['A2'] = "Khu vực khảo sát: Tuyến Metro Số 2 (Bến Thành - Tham Lương) | Nguồn dữ liệu: Sở Quy hoạch - Kiến trúc TP.HCM | Độ phân giải lưới: 4m"
ws1['A2'].font = font_subtitle
ws1['A2'].fill = fill_navy_hdr
ws1['A2'].alignment = Alignment(horizontal='center', vertical='center')
ws1.row_dimensions[2].height = 20

kpis = [
    ('TỔNG SỐ THỬA ĐẤT', f"{len(df1):,}", '100% diện tích khảo sát', 'A', 'B'),
    ('TỔNG DIỆN TÍCH', f"{df1['Diện Tích Thửa (m²)'].sum()/10000:.2f} ha", f"{df1['Diện Tích Thửa (m²)'].sum():,.1f} m²", 'C', 'D'),
    ('THỬA CÓ ĐẤT Ở (THỔ CƯ)', f"{enriched['Có Đất Ở'].value_counts().get('Có', 0):,}", f"{enriched['Có Đất Ở'].value_counts().get('Có', 0)/len(df1)*100:.1f}% tổng số thửa", 'E', 'F'),
    ('THỬA SẠCH (KO CÂY XANH)', f"{enriched['Có Cây Xanh'].value_counts().get('Không', 0):,}", f"{enriched['Có Cây Xanh'].value_counts().get('Không', 0)/len(df1)*100:.1f}% an toàn quy hoạch", 'G', 'H'),
    ('DÍNH QUY HOẠCH CÂY XANH', f"{enriched['Có Cây Xanh'].value_counts().get('CÓ (DÍNH)', 0):,}", f"{enriched['Có Cây Xanh'].value_counts().get('CÓ (DÍNH)', 0)/len(df1)*100:.1f}% cần cảnh báo", 'I', 'J'),
    ('DÍNH LỘ GIỚI / GIAO THÔNG', f"{enriched['Dính Lộ Giới / Giao Thông'].value_counts().get('Có', 0):,}", f"{enriched['Dính Lộ Giới / Giao Thông'].value_counts().get('Có', 0)/len(df1)*100:.1f}% tiếp giáp mở đường", 'K', 'L')
]

for title, val, sub, col_start, col_end in kpis:
    ws1.merge_cells(f'{col_start}4:{col_end}4')
    ws1.merge_cells(f'{col_start}5:{col_end}5')
    ws1.merge_cells(f'{col_start}6:{col_end}6')
    
    ws1[f'{col_start}4'] = title
    ws1[f'{col_start}4'].font = font_kpi_label
    ws1[f'{col_start}4'].alignment = align_center
    ws1[f'{col_start}4'].fill = fill_kpi_bg
    
    ws1[f'{col_start}5'] = val
    ws1[f'{col_start}5'].font = font_kpi_num
    ws1[f'{col_start}5'].alignment = align_center
    ws1[f'{col_start}5'].fill = fill_kpi_bg
    
    ws1[f'{col_start}6'] = sub
    ws1[f'{col_start}6'].font = font_kpi_sub
    ws1[f'{col_start}6'].alignment = align_center
    ws1[f'{col_start}6'].fill = fill_kpi_bg
    
    for r in range(4, 7):
        for c_char in [col_start, col_end]:
            ws1[f'{c_char}{r}'].border = border_cell

ws1.row_dimensions[4].height = 18
ws1.row_dimensions[5].height = 28
ws1.row_dimensions[6].height = 16

# Table 1: Phân bổ theo Quận / Huyện
ws1['A8'] = "1. BẢNG THỐNG KÊ CHI TIẾT THEO QUẬN / HUYỆN"
ws1['A8'].font = font_section

dist_agg = df1_enriched.groupby('Quận / Huyện').agg(
    Tong_Thua=('STT', 'count'),
    Tong_DT_m2=('Diện Tích Thửa (m²)', 'sum'),
    Thua_Dat_O=('Có Đất Ở', lambda x: (x == 'Có').sum()),
    Thua_Cay_Xanh=('Có Cây Xanh', lambda x: (x == 'CÓ (DÍNH)').sum()),
    Thua_Lo_Gioi=('Dính Lộ Giới / Giao Thông', lambda x: (x == 'Có').sum()),
    Thua_1500=('Dự Án 1/500 & ĐCCB', lambda x: (x != 'Không').sum())
).reset_index()

dist_agg['Ty_Le_Thua_%'] = dist_agg['Tong_Thua'] / len(df1) * 100
dist_agg['Ty_Le_Dat_O_%'] = dist_agg['Thua_Dat_O'] / dist_agg['Tong_Thua'] * 100
dist_agg['Ty_Le_CX_%'] = dist_agg['Thua_Cay_Xanh'] / dist_agg['Tong_Thua'] * 100
dist_agg = dist_agg.sort_values(by='Tong_Thua', ascending=False)

tbl1_headers = ['Quận / Huyện', 'Số Thửa Khảo Sát', 'Tỷ Lệ Thửa (%)', 'Tổng Diện Tích (m²)', 'Tổng DT (ha)', 'Thửa Có Đất Ở', 'Tỷ Lệ Đất Ở (%)', 'Thửa Dính Cây Xanh', 'Tỷ Lệ Cây Xanh (%)', 'Thửa Dính Lộ Giới', 'Thửa Dính DA 1/500 & ĐCCB']
ws1.row_dimensions[9].height = 24
for col_idx, h in enumerate(tbl1_headers, start=1):
    cell = ws1.cell(row=9, column=col_idx, value=h)
    cell.font = font_table_hdr
    cell.fill = fill_blue_hdr
    cell.alignment = align_center
    cell.border = border_header

curr_row = 10
for _, r in dist_agg.iterrows():
    ws1.row_dimensions[curr_row].height = 20
    ws1.cell(row=curr_row, column=1, value=r['Quận / Huyện']).alignment = align_left
    ws1.cell(row=curr_row, column=2, value=int(r['Tong_Thua'])).number_format = '#,##0'
    ws1.cell(row=curr_row, column=3, value=r['Ty_Le_Thua_%'] / 100).number_format = '0.00%'
    ws1.cell(row=curr_row, column=4, value=r['Tong_DT_m2']).number_format = '#,##0.0'
    ws1.cell(row=curr_row, column=5, value=r['Tong_DT_m2'] / 10000).number_format = '#,##0.00'
    ws1.cell(row=curr_row, column=6, value=int(r['Thua_Dat_O'])).number_format = '#,##0'
    ws1.cell(row=curr_row, column=7, value=r['Ty_Le_Dat_O_%'] / 100).number_format = '0.00%'
    ws1.cell(row=curr_row, column=8, value=int(r['Thua_Cay_Xanh'])).number_format = '#,##0'
    ws1.cell(row=curr_row, column=9, value=r['Ty_Le_CX_%'] / 100).number_format = '0.00%'
    ws1.cell(row=curr_row, column=10, value=int(r['Thua_Lo_Gioi'])).number_format = '#,##0'
    ws1.cell(row=curr_row, column=11, value=int(r['Thua_1500'])).number_format = '#,##0'
    
    for c in range(1, 12):
        cell = ws1.cell(row=curr_row, column=c)
        cell.font = font_cell
        cell.border = border_cell
        if c > 1:
            cell.alignment = align_right
        if curr_row % 2 == 1:
            cell.fill = fill_zebra
    curr_row += 1

# Table 1 Totals
ws1.row_dimensions[curr_row].height = 22
ws1.cell(row=curr_row, column=1, value="TỔNG CỘNG").alignment = align_left
ws1.cell(row=curr_row, column=2, value=f"=SUM(B10:B{curr_row-1})").number_format = '#,##0'
ws1.cell(row=curr_row, column=3, value=1.0).number_format = '0.00%'
ws1.cell(row=curr_row, column=4, value=f"=SUM(D10:D{curr_row-1})").number_format = '#,##0.0'
ws1.cell(row=curr_row, column=5, value=f"=SUM(E10:E{curr_row-1})").number_format = '#,##0.00'
ws1.cell(row=curr_row, column=6, value=f"=SUM(F10:F{curr_row-1})").number_format = '#,##0'
ws1.cell(row=curr_row, column=7, value=f"=F{curr_row}/B{curr_row}").number_format = '0.00%'
ws1.cell(row=curr_row, column=8, value=f"=SUM(H10:H{curr_row-1})").number_format = '#,##0'
ws1.cell(row=curr_row, column=9, value=f"=H{curr_row}/B{curr_row}").number_format = '0.00%'
ws1.cell(row=curr_row, column=10, value=f"=SUM(J10:J{curr_row-1})").number_format = '#,##0'
ws1.cell(row=curr_row, column=11, value=f"=SUM(K10:K{curr_row-1})").number_format = '#,##0'

for c in range(1, 12):
    cell = ws1.cell(row=curr_row, column=c)
    cell.font = font_cell_bold
    cell.fill = fill_total_row
    cell.border = border_total

# Table 2: Cơ cấu theo 8 Nhóm
curr_row += 3
ws1.cell(row=curr_row, column=1, value="2. BẢNG PHÂN TÍCH CƠ CẤU THEO 8 NHÓM QUY HOẠCH SỬ DỤNG ĐẤT").font = font_section

curr_row += 1
tbl2_headers = ['Nhóm Quy Hoạch Sử Dụng Đất', 'Số Ô Phân Khu', 'Số Thửa Tiếp Giáp', 'Tổng Diện Tích Ô (m²)', 'Tổng DT (ha)', 'Tỷ Lệ Diện Tích (%)', 'Mức Độ Rủi Ro / Khuyến Nghị Pháp Lý']
ws1.row_dimensions[curr_row].height = 24
for col_idx, h in enumerate(tbl2_headers, start=1):
    cell = ws1.cell(row=curr_row, column=col_idx, value=h)
    cell.font = font_table_hdr
    cell.fill = fill_teal_hdr
    cell.alignment = align_center
    cell.border = border_header

group_agg = df2.groupby('Nhom_Quy_Hoach').agg(
    So_O=('STT Thửa', 'count'),
    So_Thua=('STT Thửa', 'nunique'),
    Tong_DT_m2=('Diện Tích Ô (m²)', 'sum'),
    Muc_Do_Rui_Ro=('Muc_Do_Rui_Ro', 'first')
).reset_index().sort_values(by='Tong_DT_m2', ascending=False)

tot_dt_o = df2['Diện Tích Ô (m²)'].sum()

t2_start_row = curr_row + 1
curr_row += 1
for _, gr in group_agg.iterrows():
    ws1.row_dimensions[curr_row].height = 20
    ws1.cell(row=curr_row, column=1, value=gr['Nhom_Quy_Hoach']).alignment = align_left
    ws1.cell(row=curr_row, column=2, value=int(gr['So_O'])).number_format = '#,##0'
    ws1.cell(row=curr_row, column=3, value=int(gr['So_Thua'])).number_format = '#,##0'
    ws1.cell(row=curr_row, column=4, value=gr['Tong_DT_m2']).number_format = '#,##0.0'
    ws1.cell(row=curr_row, column=5, value=gr['Tong_DT_m2'] / 10000).number_format = '#,##0.00'
    ws1.cell(row=curr_row, column=6, value=gr['Tong_DT_m2'] / tot_dt_o).number_format = '0.00%'
    ws1.cell(row=curr_row, column=7, value=gr['Muc_Do_Rui_Ro']).alignment = align_left
    
    for c in range(1, 8):
        cell = ws1.cell(row=curr_row, column=c)
        cell.font = font_cell
        cell.border = border_cell
        if 2 <= c <= 6:
            cell.alignment = align_right
        if curr_row % 2 == 1:
            cell.fill = fill_zebra
    curr_row += 1

# Table 2 Totals
ws1.row_dimensions[curr_row].height = 22
ws1.cell(row=curr_row, column=1, value="TỔNG CỘNG CÁC Ô QUY HOẠCH").alignment = align_left
ws1.cell(row=curr_row, column=2, value=f"=SUM(B{t2_start_row}:B{curr_row-1})").number_format = '#,##0'
ws1.cell(row=curr_row, column=3, value="-").alignment = align_center
ws1.cell(row=curr_row, column=4, value=f"=SUM(D{t2_start_row}:D{curr_row-1})").number_format = '#,##0.0'
ws1.cell(row=curr_row, column=5, value=f"=SUM(E{t2_start_row}:E{curr_row-1})").number_format = '#,##0.00'
ws1.cell(row=curr_row, column=6, value=1.0).number_format = '0.00%'
ws1.cell(row=curr_row, column=7, value="").alignment = align_left

for c in range(1, 8):
    cell = ws1.cell(row=curr_row, column=c)
    cell.font = font_cell_bold
    cell.fill = fill_total_row
    cell.border = border_total

# ==========================================
# SHEET 2: DANH MỤC 72 LOẠI ĐẤT
# ==========================================
print("3. Writing Sheet 2: Danh mục 72 loại đất...")
ws2 = wb.create_sheet(title='📚 2. Danh Mục 72 Loại Đất')
ws2.views.sheetView[0].showGridLines = True

ws2.merge_cells('A1:K1')
ws2['A1'] = "TỪ ĐIỂN & DANH MỤC PHÂN LOẠI 72 LOẠI ĐẤT QUY HOẠCH ĐÔ THỊ (SQHKT TP.HCM)"
ws2['A1'].font = font_title
ws2['A1'].fill = fill_navy_hdr
ws2['A1'].alignment = align_center
ws2.row_dimensions[1].height = 30

ws2.merge_cells('A2:K2')
ws2['A2'] = "Bảng chuẩn hóa các chức năng quy hoạch phân khu 1/2000, ý nghĩa sử dụng đất, khả năng cấp phép xây dựng và mức độ rủi ro đầu tư"
ws2['A2'].font = font_subtitle
ws2['A2'].fill = fill_navy_hdr
ws2['A2'].alignment = align_center
ws2.row_dimensions[2].height = 18

s2_headers = [
    'STT', 'Tên Chức Năng SQHKT', 'Nhóm Quy Hoạch Chuẩn Hóa', 'Ý Nghĩa / Mục Đích Sử Dụng Đất',
    'Khả Năng Cấp Phép Xây Dựng', 'Đánh Giá Mức Độ Rủi Ro', 'Số Ô Xuất Hiện', 'Số Thửa Duy Nhất',
    'Tổng DT Ô (m²)', 'Mật Độ XD Max (%)', 'HSSDĐ Max'
]
ws2.row_dimensions[3].height = 26
for c_idx, h in enumerate(s2_headers, start=1):
    cell = ws2.cell(row=3, column=c_idx, value=h)
    cell.font = font_table_hdr
    cell.fill = fill_blue_hdr
    cell.alignment = align_center
    cell.border = border_header

for idx, r in cat_df.iterrows():
    row_num = idx + 4
    ws2.row_dimensions[row_num].height = 20
    ws2.cell(row=row_num, column=1, value=r['STT']).alignment = align_center
    ws2.cell(row=row_num, column=2, value=r['Tên Chức Năng SQHKT']).alignment = align_left
    ws2.cell(row=row_num, column=3, value=r['Nhom_Quy_Hoach']).alignment = align_left
    ws2.cell(row=row_num, column=4, value=r['Y_Nghia']).alignment = align_left
    ws2.cell(row=row_num, column=5, value=r['Kha_Nang_XD']).alignment = align_left
    ws2.cell(row=row_num, column=6, value=r['Muc_Do_Rui_Ro']).alignment = align_left
    ws2.cell(row=row_num, column=7, value=int(r['So_Lan_Xuat_Hien'])).number_format = '#,##0'
    ws2.cell(row=row_num, column=8, value=int(r['So_Thua_Duy_Nhat'])).number_format = '#,##0'
    ws2.cell(row=row_num, column=9, value=r['Tong_DT_O_m2']).number_format = '#,##0.0'
    ws2.cell(row=row_num, column=10, value=r['Mat_Do_Max'] if pd.notna(r['Mat_Do_Max']) else '-').alignment = align_right
    ws2.cell(row=row_num, column=11, value=r['HSSDD_Max'] if pd.notna(r['HSSDD_Max']) else '-').alignment = align_right
    
    for c in range(1, 12):
        cell = ws2.cell(row=row_num, column=c)
        cell.font = font_cell
        cell.border = border_cell
        if 7 <= c <= 9:
            cell.alignment = align_right
        if row_num % 2 == 1:
            cell.fill = fill_zebra
        if 'RỦI RO CAO' in str(r['Muc_Do_Rui_Ro']):
            if c == 6:
                cell.fill = fill_highlight_red
        elif 'AN TOÀN' in str(r['Muc_Do_Rui_Ro']):
            if c == 6:
                cell.fill = fill_highlight_green

ws2.auto_filter.ref = f"A3:K{len(cat_df)+3}"
ws2.freeze_panes = 'C4'

# ==========================================
# SHEET 3: DANH SÁCH THỔ CƯ - ĐẤT Ở
# ==========================================
print("4. Writing Sheet 3: Danh sách Thổ cư - Đất ở...")
ws3 = wb.create_sheet(title='🏡 3. Thổ Cư - Đất Ở (4798 Thửa)')
ws3.views.sheetView[0].showGridLines = True

ws3.merge_cells('A1:N1')
ws3['A1'] = "DANH SÁCH TOÀN BỘ CÁC THỬA ĐẤT CÓ CHỨC NĂNG ĐẤT Ở / THỔ CƯ (4,798 THỬA)"
ws3['A1'].font = font_title
ws3['A1'].fill = fill_navy_hdr
ws3['A1'].alignment = align_center
ws3.row_dimensions[1].height = 28

df_tho_cu = df1_enriched[df1_enriched['Có Đất Ở'] == 'Có'].copy()

s3_headers = [
    'STT', 'Mã Thửa Đất', 'Số Tờ', 'Số Thửa', 'Quận / Huyện', 'Phường / Xã',
    'Tổng DT Thửa (m²)', 'Tỷ Lệ Đất Ở (%)', 'DT Đất Ở (m²)', 'Cảnh Báo Cây Xanh',
    'Dính Lộ Giới', 'Dự Án 1/500 & ĐCCB', 'Đồ Án 1/2000', 'Toạ Độ Khảo Sát'
]
ws3.row_dimensions[2].height = 24
for c_idx, h in enumerate(s3_headers, start=1):
    cell = ws3.cell(row=2, column=c_idx, value=h)
    cell.font = font_table_hdr
    cell.fill = fill_blue_hdr
    cell.alignment = align_center
    cell.border = border_header

for idx, (_, r) in enumerate(df_tho_cu.iterrows(), start=1):
    row_num = idx + 2
    ws3.row_dimensions[row_num].height = 18
    ws3.cell(row=row_num, column=1, value=idx).alignment = align_center
    ws3.cell(row=row_num, column=2, value=str(r['Mã Thửa Đất'])).alignment = align_left
    ws3.cell(row=row_num, column=3, value=r['Số Tờ']).alignment = align_center
    ws3.cell(row=row_num, column=4, value=r['Số Thửa']).alignment = align_center
    ws3.cell(row=row_num, column=5, value=r['Quận / Huyện']).alignment = align_left
    ws3.cell(row=row_num, column=6, value=r['Phường / Xã']).alignment = align_left
    ws3.cell(row=row_num, column=7, value=r['Diện Tích Thửa (m²)']).number_format = '#,##0.0'
    ws3.cell(row=row_num, column=8, value=r['Tỷ Lệ Đất Ở (%)'] / 100).number_format = '0.0%'
    ws3.cell(row=row_num, column=9, value=r['Diện Tích Đất Ở (m²)']).number_format = '#,##0.0'
    ws3.cell(row=row_num, column=10, value=r['Có Cây Xanh']).alignment = align_center
    ws3.cell(row=row_num, column=11, value=r['Dính Lộ Giới / Giao Thông']).alignment = align_center
    ws3.cell(row=row_num, column=12, value=r['Dự Án 1/500 & ĐCCB']).alignment = align_center
    ws3.cell(row=row_num, column=13, value=r['Đồ Án Quy Hoạch 1/2000']).alignment = align_left
    ws3.cell(row=row_num, column=14, value=r['Toạ Độ Khảo Sát (Kinh độ, Vĩ độ)']).alignment = align_left
    
    for c in range(1, 15):
        cell = ws3.cell(row=row_num, column=c)
        cell.font = font_cell
        cell.border = border_cell
        if c in [7, 8, 9]:
            cell.alignment = align_right
        if row_num % 2 == 1:
            cell.fill = fill_zebra
        if c == 10 and 'DÍNH' in str(r['Có Cây Xanh']):
            cell.fill = fill_highlight_red
            cell.font = font_cell_bold

ws3.auto_filter.ref = f"A2:N{len(df_tho_cu)+2}"
ws3.freeze_panes = 'C3'

# ==========================================
# SHEET 4: ĐẤT SẠCH (KHÔNG DÍNH CÂY XANH)
# ==========================================
print("5. Writing Sheet 4: Đất sạch không dính cây xanh...")
ws4 = wb.create_sheet(title='🌿 4. Đất Sạch (Ko Cây Xanh)')
ws4.views.sheetView[0].showGridLines = True

ws4.merge_cells('A1:M1')
ws4['A1'] = "DANH SÁCH CÁC THỬA ĐẤT AN TOÀN - HOÀN TOÀN KHÔNG DÍNH QUY HOẠCH CÂY XANH (6,139 THỬA)"
ws4['A1'].font = font_title
ws4['A1'].fill = fill_navy_hdr
ws4['A1'].alignment = align_center
ws4.row_dimensions[1].height = 28

df_sach = df1_enriched[df1_enriched['Có Cây Xanh'] == 'Không'].copy()

s4_headers = [
    'STT', 'Mã Thửa Đất', 'Số Tờ', 'Số Thửa', 'Quận / Huyện', 'Phường / Xã',
    'Tổng DT Thửa (m²)', 'Phân Loại Quy Hoạch', 'Có Đất Ở', 'Tỷ Lệ Đất Ở (%)',
    'Dính Lộ Giới', 'Dự Án 1/500 & ĐCCB', 'Cơ Cấu Quy Hoạch Sử Dụng Đất'
]
ws4.row_dimensions[2].height = 24
for c_idx, h in enumerate(s4_headers, start=1):
    cell = ws4.cell(row=2, column=c_idx, value=h)
    cell.font = font_table_hdr
    cell.fill = fill_teal_hdr
    cell.alignment = align_center
    cell.border = border_header

for idx, (_, r) in enumerate(df_sach.iterrows(), start=1):
    row_num = idx + 2
    ws4.row_dimensions[row_num].height = 18
    ws4.cell(row=row_num, column=1, value=idx).alignment = align_center
    ws4.cell(row=row_num, column=2, value=str(r['Mã Thửa Đất'])).alignment = align_left
    ws4.cell(row=row_num, column=3, value=r['Số Tờ']).alignment = align_center
    ws4.cell(row=row_num, column=4, value=r['Số Thửa']).alignment = align_center
    ws4.cell(row=row_num, column=5, value=r['Quận / Huyện']).alignment = align_left
    ws4.cell(row=row_num, column=6, value=r['Phường / Xã']).alignment = align_left
    ws4.cell(row=row_num, column=7, value=r['Diện Tích Thửa (m²)']).number_format = '#,##0.0'
    ws4.cell(row=row_num, column=8, value=r['Phân Loại Quy Hoạch']).alignment = align_left
    ws4.cell(row=row_num, column=9, value=r['Có Đất Ở']).alignment = align_center
    ws4.cell(row=row_num, column=10, value=r['Tỷ Lệ Đất Ở (%)'] / 100).number_format = '0.0%'
    ws4.cell(row=row_num, column=11, value=r['Dính Lộ Giới / Giao Thông']).alignment = align_center
    ws4.cell(row=row_num, column=12, value=r['Dự Án 1/500 & ĐCCB']).alignment = align_center
    ws4.cell(row=row_num, column=13, value=r['Cơ Cấu Quy Hoạch Sử Dụng Đất']).alignment = align_left
    
    for c in range(1, 14):
        cell = ws4.cell(row=row_num, column=c)
        cell.font = font_cell
        cell.border = border_cell
        if c in [7, 10]:
            cell.alignment = align_right
        if row_num % 2 == 1:
            cell.fill = fill_zebra

ws4.auto_filter.ref = f"A2:M{len(df_sach)+2}"
ws4.freeze_panes = 'C3'

# ==========================================
# SHEET 5: CẢNH BÁO DÍNH CÂY XANH
# ==========================================
print("6. Writing Sheet 5: Cảnh báo dính cây xanh...")
ws5 = wb.create_sheet(title='⚠️ 5. Cảnh Báo Dính Cây Xanh')
ws5.views.sheetView[0].showGridLines = True

ws5.merge_cells('A1:L1')
ws5['A1'] = "DANH SÁCH THỬA ĐẤT BỊ DÍNH QUY HOẠCH CÂY XANH, CÔNG VIÊN & CÁCH LY (182 THỬA)"
ws5['A1'].font = font_title
ws5['A1'].fill = PatternFill(start_color='991B1B', end_color='991B1B', fill_type='solid')
ws5['A1'].alignment = align_center
ws5.row_dimensions[1].height = 28

df_cx = df1_enriched[df1_enriched['Có Cây Xanh'] == 'CÓ (DÍNH)'].sort_values(by='Tỷ Lệ Cây Xanh (%)', ascending=False).copy()

s5_headers = [
    'STT', 'Mã Thửa Đất', 'Số Tờ', 'Số Thửa', 'Quận / Huyện', 'Phường / Xã',
    'Tổng DT Thửa (m²)', 'Loại QH Cây Xanh Dính Phải', 'Tỷ Lệ Dính CX (%)',
    'Diện Tích Dính CX (m²)', 'Có Đất Ở Còn Lại', 'Tỷ Lệ Đất Ở (%)'
]
ws5.row_dimensions[2].height = 24
for c_idx, h in enumerate(s5_headers, start=1):
    cell = ws5.cell(row=2, column=c_idx, value=h)
    cell.font = font_table_hdr
    cell.fill = PatternFill(start_color='B91C1C', end_color='B91C1C', fill_type='solid')
    cell.alignment = align_center
    cell.border = border_header

for idx, (_, r) in enumerate(df_cx.iterrows(), start=1):
    row_num = idx + 2
    ws5.row_dimensions[row_num].height = 18
    ws5.cell(row=row_num, column=1, value=idx).alignment = align_center
    ws5.cell(row=row_num, column=2, value=str(r['Mã Thửa Đất'])).alignment = align_left
    ws5.cell(row=row_num, column=3, value=r['Số Tờ']).alignment = align_center
    ws5.cell(row=row_num, column=4, value=r['Số Thửa']).alignment = align_center
    ws5.cell(row=row_num, column=5, value=r['Quận / Huyện']).alignment = align_left
    ws5.cell(row=row_num, column=6, value=r['Phường / Xã']).alignment = align_left
    ws5.cell(row=row_num, column=7, value=r['Diện Tích Thửa (m²)']).number_format = '#,##0.0'
    ws5.cell(row=row_num, column=8, value=r['Loại QH Cây Xanh Dính Phải']).alignment = align_left
    ws5.cell(row=row_num, column=9, value=r['Tỷ Lệ Cây Xanh (%)'] / 100).number_format = '0.0%'
    ws5.cell(row=row_num, column=10, value=r['Diện Tích Cây Xanh (m²)']).number_format = '#,##0.0'
    ws5.cell(row=row_num, column=11, value=r['Có Đất Ở']).alignment = align_center
    ws5.cell(row=row_num, column=12, value=r['Tỷ Lệ Đất Ở (%)'] / 100).number_format = '0.0%'
    
    for c in range(1, 13):
        cell = ws5.cell(row=row_num, column=c)
        cell.font = font_cell
        cell.border = border_cell
        if c in [7, 9, 10, 12]:
            cell.alignment = align_right
        if row_num % 2 == 1:
            cell.fill = fill_highlight_red

ws5.auto_filter.ref = f"A2:L{len(df_cx)+2}"
ws5.freeze_panes = 'C3'

# ==========================================
# SHEET 6: DỰ ÁN 1-500 & ĐIỀU CHỈNH CỤC BỘ (632 BẢN GHI)
# ==========================================
print("7. Writing Sheet 6: Dự án 1/500 & điều chỉnh cục bộ từ Sheet 4...")
ws6 = wb.create_sheet(title='🏗️ 6. Dự Án 1-500 & Điều Chỉnh')
ws6.views.sheetView[0].showGridLines = True

ws6.merge_cells('A1:L1')
ws6['A1'] = "CHI TIẾT CÁC THỬA ĐẤT THUỘC ĐỒ ÁN QUY HOẠCH 1/500 & QUYẾT ĐỊNH ĐIỀU CHỈNH CỤC BỘ (632 BẢN GHI)"
ws6['A1'].font = font_title
ws6['A1'].fill = fill_navy_hdr
ws6['A1'].alignment = align_center
ws6.row_dimensions[1].height = 28

s6_headers = [
    'STT', 'Mã Thửa Đất', 'Số Tờ', 'Số Thửa', 'Quận / Huyện', 'Phường / Xã',
    'Loại Quy Hoạch', 'Tên Dự Án / Quyết Định Điều Chỉnh', 'Số Quyết Định',
    'Ngày Phê Duyệt', 'Cơ Quan Phê Duyệt', 'Tỷ Lệ Thửa Thuộc Dự Án (%)'
]
ws6.row_dimensions[2].height = 24
for c_idx, h in enumerate(s6_headers, start=1):
    cell = ws6.cell(row=2, column=c_idx, value=h)
    cell.font = font_table_hdr
    cell.fill = fill_purple_hdr
    cell.alignment = align_center
    cell.border = border_header

for idx, (_, r) in enumerate(df4.iterrows(), start=1):
    row_num = idx + 2
    ws6.row_dimensions[row_num].height = 18
    ws6.cell(row=row_num, column=1, value=idx).alignment = align_center
    ws6.cell(row=row_num, column=2, value=str(r['Mã Thửa Đất'])).alignment = align_left
    ws6.cell(row=row_num, column=3, value=r['Số Tờ']).alignment = align_center
    ws6.cell(row=row_num, column=4, value=r['Số Thửa']).alignment = align_center
    ws6.cell(row=row_num, column=5, value=r['Quận / Huyện']).alignment = align_left
    ws6.cell(row=row_num, column=6, value=r['Phường / Xã']).alignment = align_left
    ws6.cell(row=row_num, column=7, value=r['Loại Quy Hoạch']).alignment = align_center
    ws6.cell(row=row_num, column=8, value=r['Tên Dự Án / Đồ Án Điều Chỉnh']).alignment = align_left
    ws6.cell(row=row_num, column=9, value=r['Số Quyết Định Phê Duyệt'] if pd.notna(r['Số Quyết Định Phê Duyệt']) else '-').alignment = align_left
    ws6.cell(row=row_num, column=10, value=str(r['Ngày Phê Duyệt']) if pd.notna(r['Ngày Phê Duyệt']) else '-').alignment = align_center
    ws6.cell(row=row_num, column=11, value=r['Cơ Quan Phê Duyệt'] if pd.notna(r['Cơ Quan Phê Duyệt']) else '-').alignment = align_left
    
    pct_val = r['Tỷ Lệ Thửa Thuộc Dự Án (%)']
    if pd.notna(pct_val):
        ws6.cell(row=row_num, column=12, value=float(pct_val) if float(pct_val) <= 1.0 else float(pct_val)/100).number_format = '0.0%'
    else:
        ws6.cell(row=row_num, column=12, value='-').alignment = align_center
    
    for c in range(1, 13):
        cell = ws6.cell(row=row_num, column=c)
        cell.font = font_cell
        cell.border = border_cell
        if c == 12:
            cell.alignment = align_right
        if row_num % 2 == 1:
            cell.fill = fill_zebra
        if '1/500' in str(r['Loại Quy Hoạch']):
            if c == 7:
                cell.fill = fill_highlight_amber

ws6.auto_filter.ref = f"A2:L{len(df4)+2}"
ws6.freeze_panes = 'C3'

# Auto-adjust column widths for all sheets
print("8. Auto-adjusting column widths...")
for sheet in wb.worksheets:
    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in [1, 2] and col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N']:
                continue
            if cell.value:
                val_str = str(cell.value)
                max_len = max(max_len, len(val_str))
        sheet.column_dimensions[col_letter].width = max(max_len + 3, 11)

ws1.column_dimensions['A'].width = 30
ws1.column_dimensions['B'].width = 18
ws1.column_dimensions['C'].width = 18
ws1.column_dimensions['D'].width = 22
ws1.column_dimensions['E'].width = 16
ws1.column_dimensions['F'].width = 18
ws1.column_dimensions['G'].width = 38

print(f"9. Saving workbook to: {output_excel}")
wb.save(output_excel)
print("SUCCESS: Excel report created successfully!")
