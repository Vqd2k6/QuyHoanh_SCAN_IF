import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import numpy as np
import unicodedata

excel_path = 'data/output/KS002/KS002.xlsx'
output_excel = 'data/output/KS002/Bao_Cao_Phan_Tich_Quy_Hoach_KS002.xlsx'

print("1. Loading raw data from:", excel_path)
df1 = pd.read_excel(excel_path, sheet_name=0, header=3)
df2 = pd.read_excel(excel_path, sheet_name=1, header=3)
df3 = pd.read_excel(excel_path, sheet_name=2, header=3)
df4 = pd.read_excel(excel_path, sheet_name=3, header=3)

def normalize(s):
    if not isinstance(s, str):
        return ''
    return unicodedata.normalize('NFC', s).strip()

df2['norm_land'] = df2['Chức Năng Sử Dụng Đất'].apply(normalize)
df2['norm_detail'] = df2['Chức Năng Chi Tiết (SQHKT)'].apply(normalize)
df2['norm_block'] = df2['Mã Ô Phố'].apply(lambda x: '' if str(x).strip() in ['-', '', 'None', 'nan'] else f"[{str(x).strip()}] ")

# Create full detailed label: e.g. "[II.26] Cây xanh cách ly: 248.47 m² (3.1%)"
def build_label(r):
    code = r['norm_block']
    name = r['norm_land']
    dt = r['Diện Tích Ô (m²)']
    pct = r['Tỷ Lệ Chiếm Thửa (%)'] * 100
    return f"{code}{name}: {dt:,.2f} m² ({pct:.1f}%)"

df2['Nhãn Ô Chi Tiết'] = df2.apply(build_label, axis=1)

def get_category_broad(name):
    n = name.lower()
    if any(k in n for k in ['cây xanh', 'công viên', 'cv -', 'cv-', 'thể dục thể thao']):
        return '5. Cây Xanh, Công Viên & TDTT'
    if any(k in n for k in ['đất ở', 'dân cư', 'nhà ở', 'nhóm nhà ở']):
        return '1. Đất Ở / Thổ Cư & Dân Cư'
    if any(k in n for k in ['phức hợp', 'hỗn hợp']):
        return '2. Đất Phức Hợp / Hỗn Hợp'
    if any(k in n for k in ['thương mại', 'tmdv', 'sxkd', 'cơ quan']):
        return '3. Thương Mại - Dịch Vụ & Cơ Quan'
    if any(k in n for k in ['công cộng', 'giáo dục', 'trường học', 'y tế', 'hành chính', 'hành chánh', 'tôn giáo', 'di tích', 'quảng trường']):
        return '4. Công Trình Công Cộng & An Sinh'
    if any(k in n for k in ['giao thông', 'ga depot', 'depot', 'mặt nước', 'kênh', 'rạch', 'sông', 'hạ tầng', 'cấp nước', 'kho tàng']):
        return '6. Giao Thông, Hạ Tầng & Mặt Nước'
    if any(k in n for k in ['công nghiệp', 'tiểu thủ', 'quân sự', 'quốc phòng']):
        return '7. Công Nghiệp & Quốc Phòng'
    return '8. Đất Khác'

df2['Nhóm Quy Hoạch Tham Khảo'] = df2['norm_land'].apply(get_category_broad)

tot_survey_area = df1['Diện Tích Thửa (m²)'].sum()
tot_zone_area = df2['Diện Tích Ô (m²)'].sum()

raw_types_df = df2.groupby('norm_land').agg(
    Nhóm_Tham_Khảo=('Nhóm Quy Hoạch Tham Khảo', 'first'),
    Số_Ô=('STT Thửa', 'count'),
    Số_Thửa_Tiếp_Giáp=('STT Thửa', 'nunique'),
    Tổng_DT_Ô_m2=('Diện Tích Ô (m²)', 'sum'),
    DT_Trung_Bình_Ô_m2=('Diện Tích Ô (m²)', 'mean'),
    Tầng_Cao_Max=('Tầng Cao Cho Phép (tầng)', lambda x: pd.to_numeric(x, errors='coerce').max()),
    Mật_Độ_Max=('Mật Độ XD Tối Đa (%)', lambda x: pd.to_numeric(x, errors='coerce').max()),
    HSSDĐ_Max=('Hệ Số Sử Dụng Đất (HSSDĐ)', lambda x: pd.to_numeric(x, errors='coerce').max()),
    Mã_Ô_Phố_Tiêu_Biểu=('Mã Ô Phố', lambda x: ', '.join([str(i) for i in x.unique() if str(i).strip() not in ['-', '', 'None', 'nan']][:6]))
).reset_index().rename(columns={'norm_land': 'Tên Loại Đất Thô (SQHKT)'})

raw_types_df['Tỷ Lệ DT Trên Tổng Khảo Sát (%)'] = raw_types_df['Tổng_DT_Ô_m2'] / tot_survey_area * 100
raw_types_df = raw_types_df.sort_values(by='Tổng_DT_Ô_m2', ascending=False).reset_index(drop=True)
raw_types_df.insert(0, 'STT', range(1, len(raw_types_df) + 1))

res_agg = df2[df2['Nhóm Quy Hoạch Tham Khảo'] == '1. Đất Ở / Thổ Cư & Dân Cư'].groupby('STT Thửa').agg(
    DT_Dat_O=('Diện Tích Ô (m²)', 'sum'),
    Ty_Le_Dat_O=('Tỷ Lệ Chiếm Thửa (%)', 'sum'),
    Loai_Dat_O=('norm_land', lambda x: ', '.join(x.unique()))
).reset_index()
res_dict = {r['STT Thửa']: r for _, r in res_agg.iterrows()}

green_agg = df2[df2['Nhóm Quy Hoạch Tham Khảo'] == '5. Cây Xanh, Công Viên & TDTT'].groupby('STT Thửa').agg(
    DT_Cay_Xanh=('Diện Tích Ô (m²)', 'sum'),
    Ty_Le_Cay_Xanh=('Tỷ Lệ Chiếm Thửa (%)', 'sum'),
    Loai_Cay_Xanh=('norm_land', lambda x: ', '.join(x.unique())),
    Nhan_Cay_Xanh=('Nhãn Ô Chi Tiết', lambda x: ' | '.join(x.unique()))
).reset_index()
green_dict = {r['STT Thửa']: r for _, r in green_agg.iterrows()}

traffic_agg = df2[df2['Nhóm Quy Hoạch Tham Khảo'] == '6. Giao Thông, Hạ Tầng & Mặt Nước'].groupby('STT Thửa').agg(
    DT_Giao_Thong=('Diện Tích Ô (m²)', 'sum'),
    Ty_Le_Giao_Thong=('Tỷ Lệ Chiếm Thửa (%)', 'sum')
).reset_index()
traffic_dict = {r['STT Thửa']: r for _, r in traffic_agg.iterrows()}

mixed_agg = df2[df2['Nhóm Quy Hoạch Tham Khảo'] == '2. Đất Phức Hợp / Hỗn Hợp'].groupby('STT Thửa').agg(
    DT_Phuc_Hop=('Diện Tích Ô (m²)', 'sum'),
    Ty_Le_Phuc_Hop=('Tỷ Lệ Chiếm Thửa (%)', 'sum')
).reset_index()
mixed_dict = {r['STT Thửa']: r for _, r in mixed_agg.iterrows()}

public_agg = df2[df2['Nhóm Quy Hoạch Tham Khảo'] == '4. Công Trình Công Cộng & An Sinh'].groupby('STT Thửa').agg(
    DT_Cong_Cong=('Diện Tích Ô (m²)', 'sum'),
    Ty_Le_Cong_Cong=('Tỷ Lệ Chiếm Thửa (%)', 'sum')
).reset_index()
public_dict = {r['STT Thửa']: r for _, r in public_agg.iterrows()}

parcels_with_1500 = set(df4[df4['Loại Quy Hoạch'] == 'Quy hoạch chi tiết 1/500']['STT Thửa'].unique())
parcels_with_dccb = set(df4[df4['Loại Quy Hoạch'] == 'Điều chỉnh cục bộ (DCCB)']['STT Thửa'].unique())

def enrich_parcel_full(row):
    stt = row['STT']
    res = res_dict.get(stt, None)
    green = green_dict.get(stt, None)
    traffic = traffic_dict.get(stt, None)
    mixed = mixed_dict.get(stt, None)
    public = public_dict.get(stt, None)
    
    dt_o = round(res['DT_Dat_O'], 2) if res is not None else 0.0
    pct_o = round(res['Ty_Le_Dat_O'] * 100, 2) if res is not None else 0.0
    loai_o = res['Loai_Dat_O'] if res is not None else 'Không'
    
    dt_cx = round(green['DT_Cay_Xanh'], 2) if green is not None else 0.0
    pct_cx = round(green['Ty_Le_Cay_Xanh'] * 100, 2) if green is not None else 0.0
    loai_cx = green['Loai_Cay_Xanh'] if green is not None else 'Không'
    nhan_cx = green['Nhan_Cay_Xanh'] if green is not None else 'Không'
    
    dt_gt = round(traffic['DT_Giao_Thong'], 2) if traffic is not None else 0.0
    pct_gt = round(traffic['Ty_Le_Giao_Thong'] * 100, 2) if traffic is not None else 0.0
    
    dt_ph = round(mixed['DT_Phuc_Hop'], 2) if mixed is not None else 0.0
    pct_ph = round(mixed['Ty_Le_Phuc_Hop'] * 100, 2) if mixed is not None else 0.0
    
    dt_cc = round(public['DT_Cong_Cong'], 2) if public is not None else 0.0
    pct_cc = round(public['Ty_Le_Cong_Cong'] * 100, 2) if public is not None else 0.0
    
    status_1500 = 'Có QH 1/500' if stt in parcels_with_1500 else ('Có ĐCCB' if stt in parcels_with_dccb else 'Không')
    
    if dt_cx > 0:
        safe_badge = '🔴 Dính Cây Xanh'
    elif dt_cc > 0:
        safe_badge = '🔴 Dính Công Cộng'
    elif dt_o > 0 and dt_cx == 0 and dt_cc == 0:
        safe_badge = '🟢 Thổ Cư An Toàn'
    elif dt_ph > 0:
        safe_badge = '🔵 Phức Hợp'
    else:
        safe_badge = '⚪ Khác'
        
    return pd.Series({
        'Đánh Giá An Toàn': safe_badge,
        'DT Đất Ở (m²)': dt_o,
        'Tỷ Lệ Đất Ở (%)': pct_o,
        'Loại Đất Ở Thô': loai_o,
        'DT Cây Xanh (m²)': dt_cx,
        'Tỷ Lệ Cây Xanh (%)': pct_cx,
        'Loại Cây Xanh Thô': loai_cx,
        'Chi Tiết Ô Cây Xanh': nhan_cx,
        'DT Giao Thông Lộ Giới (m²)': dt_gt,
        'Tỷ Lệ Giao Thông (%)': pct_gt,
        'DT Phức Hợp (m²)': dt_ph,
        'Tỷ Lệ Phức Hợp (%)': pct_ph,
        'DT Công Trình CC (m²)': dt_cc,
        'Tỷ Lệ Công Cộng (%)': pct_cc,
        'Dự Án 1/500 & ĐCCB': status_1500
    })

df1_full_enriched = pd.concat([df1, df1.apply(enrich_parcel_full, axis=1)], axis=1)

wb = openpyxl.Workbook()
wb.remove(wb.active)

FONT_NAME = 'Arial'
font_title = Font(name=FONT_NAME, size=13, bold=True, color='FFFFFF')
font_subtitle = Font(name=FONT_NAME, size=9, italic=True, color='E2E8F0')
font_section = Font(name=FONT_NAME, size=11, bold=True, color='1E3A8A')
font_table_hdr = Font(name=FONT_NAME, size=9, bold=True, color='FFFFFF')
font_kpi_num = Font(name=FONT_NAME, size=15, bold=True, color='1E3A8A')
font_kpi_label = Font(name=FONT_NAME, size=9, bold=True, color='475569')
font_kpi_sub = Font(name=FONT_NAME, size=8, italic=True, color='64748B')
font_cell = Font(name=FONT_NAME, size=9)
font_cell_bold = Font(name=FONT_NAME, size=9, bold=True)

fill_navy_hdr = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
fill_blue_hdr = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
fill_teal_hdr = PatternFill(start_color='0D9488', end_color='0D9488', fill_type='solid')
fill_purple_hdr = PatternFill(start_color='4338CA', end_color='4338CA', fill_type='solid')
fill_kpi_bg = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
fill_highlight_green = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')
fill_highlight_red = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
fill_highlight_amber = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
fill_total_row = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')

thin_side = Side(border_style='thin', color='CBD5E1')
thick_bottom = Side(border_style='medium', color='1E3A8A')
double_bottom = Side(border_style='double', color='1E3A8A')

border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
border_header = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thick_bottom)
border_total = Border(left=thin_side, right=thin_side, top=thin_side, bottom=double_bottom)

align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
align_right = Alignment(horizontal='right', vertical='center')

# SHEET 1: DASHBOARD
ws1 = wb.create_sheet(title='📊 1. Dashboard & KPIs')
ws1.views.sheetView[0].showGridLines = True

ws1.merge_cells('A1:L1')
ws1['A1'] = "BÁO CÁO PHÂN TÍCH TỔNG HỢP QUY HOẠCH ĐÔ THỊ - KHẢO SÁT KS002 (METRO TUYẾN SỐ 2)"
ws1['A1'].font = font_title
ws1['A1'].fill = fill_navy_hdr
ws1['A1'].alignment = align_center
ws1.row_dimensions[1].height = 30

ws1.merge_cells('A2:L2')
ws1['A2'] = "Độ phân giải: 4m | Tổng số điểm quét: 82.824 | 6.321 thửa đất | 10.932 ô quy hoạch phân khu 1/2000 | Nguồn: SQHKT TP.HCM"
ws1['A2'].font = font_subtitle
ws1['A2'].fill = fill_navy_hdr
ws1['A2'].alignment = align_center
ws1.row_dimensions[2].height = 18

kpis = [
    ('TỔNG SỐ THỬA ĐẤT', f"{len(df1):,}", '100% phạm vi khảo sát', 'A', 'B'),
    ('TỔNG DIỆN TÍCH', f"{tot_survey_area/10000:.2f} ha", f"{tot_survey_area:,.1f} m²", 'C', 'D'),
    ('TỔNG Ô PHÂN KHU', f"{len(df2):,}", '72 loại mục đích đất thô', 'E', 'F'),
    ('THỬA CÓ ĐẤT Ở (THỔ CƯ)', f"{len(res_agg):,}", f"{len(res_agg)/len(df1)*100:.1f}% tổng số thửa", 'G', 'H'),
    ('THỬA SẠCH (KO CÂY XANH)', f"{len(df1)-len(green_agg):,}", f"{(len(df1)-len(green_agg))/len(df1)*100:.1f}% không dính công viên", 'I', 'J'),
    ('DÍNH QUY HOẠCH CÂY XANH', f"{len(green_agg):,}", f"{len(green_agg)/len(df1)*100:.1f}% cần cảnh báo rủi ro", 'K', 'L')
]

for title, val, sub, col_start, col_end in kpis:
    ws1.merge_cells(f'{col_start}4:{col_end}4')
    ws1.merge_cells(f'{col_start}5:{col_end}5')
    ws1.merge_cells(f'{col_start}6:{col_end}6')
    
    ws1[f'{col_start}4'] = title
    ws1[f'{col_start}4'].font = font_kpi_label
    ws1[f'{col_start}4'].alignment = align_center
    ws1[f'{col_start}4'].fill = fill_kpi_bg
    
    ws1[f'{col_start}5'] = val
    ws1[f'{col_start}5'].font = font_kpi_num
    ws1[f'{col_start}5'].alignment = align_center
    ws1[f'{col_start}5'].fill = fill_kpi_bg
    
    ws1[f'{col_start}6'] = sub
    ws1[f'{col_start}6'].font = font_kpi_sub
    ws1[f'{col_start}6'].alignment = align_center
    ws1[f'{col_start}6'].fill = fill_kpi_bg
    
    for r in range(4, 7):
        for c_char in [col_start, col_end]:
            ws1[f'{c_char}{r}'].border = border_cell

ws1.row_dimensions[4].height = 18
ws1.row_dimensions[5].height = 26
ws1.row_dimensions[6].height = 16

ws1['A8'] = "1. BẢNG PHÂN BỔ THEO ĐỊA BÀN 6 QUẬN / HUYỆN"
ws1['A8'].font = font_section

dist_agg = df1_full_enriched.groupby('Quận / Huyện').agg(
    Tong_Thua=('STT', 'count'),
    Tong_DT_m2=('Diện Tích Thửa (m²)', 'sum'),
    Thua_Dat_O=('DT Đất Ở (m²)', lambda x: (x > 0).sum()),
    Thua_Cay_Xanh=('DT Cây Xanh (m²)', lambda x: (x > 0).sum()),
    Thua_Lo_Gioi=('DT Giao Thông Lộ Giới (m²)', lambda x: (x > 0).sum()),
    Thua_1500=('Dự Án 1/500 & ĐCCB', lambda x: (x != 'Không').sum())
).reset_index()

dist_agg['Ty_Le_Thua_%'] = dist_agg['Tong_Thua'] / len(df1) * 100
dist_agg['Ty_Le_Dat_O_%'] = dist_agg['Thua_Dat_O'] / dist_agg['Tong_Thua'] * 100
dist_agg['Ty_Le_CX_%'] = dist_agg['Thua_Cay_Xanh'] / dist_agg['Tong_Thua'] * 100
dist_agg = dist_agg.sort_values(by='Tong_Thua', ascending=False)

tbl1_headers = ['Quận / Huyện', 'Số Thửa Khảo Sát', 'Tỷ Lệ Thửa (%)', 'Tổng Diện Tích (m²)', 'Tổng DT (ha)', 'Thửa Có Đất Ở', 'Tỷ Lệ Đất Ở (%)', 'Thửa Dính Cây Xanh', 'Tỷ Lệ Cây Xanh (%)', 'Thửa Dính Lộ Giới', 'Thửa Có DA 1/500 & ĐCCB']
ws1.row_dimensions[9].height = 24
for col_idx, h in enumerate(tbl1_headers, start=1):
    cell = ws1.cell(row=9, column=col_idx, value=h)
    cell.font = font_table_hdr
    cell.fill = fill_blue_hdr
    cell.alignment = align_center
    cell.border = border_header

curr_row = 10
for _, r in dist_agg.iterrows():
    ws1.row_dimensions[curr_row].height = 20
    ws1.cell(row=curr_row, column=1, value=r['Quận / Huyện']).alignment = align_left
    ws1.cell(row=curr_row, column=2, value=int(r['Tong_Thua'])).number_format = '#,##0'
    ws1.cell(row=curr_row, column=3, value=r['Ty_Le_Thua_%'] / 100).number_format = '0.00%'
    ws1.cell(row=curr_row, column=4, value=r['Tong_DT_m2']).number_format = '#,##0.0'
    ws1.cell(row=curr_row, column=5, value=r['Tong_DT_m2'] / 10000).number_format = '#,##0.00'
    ws1.cell(row=curr_row, column=6, value=int(r['Thua_Dat_O'])).number_format = '#,##0'
    ws1.cell(row=curr_row, column=7, value=r['Ty_Le_Dat_O_%'] / 100).number_format = '0.00%'
    ws1.cell(row=curr_row, column=8, value=int(r['Thua_Cay_Xanh'])).number_format = '#,##0'
    ws1.cell(row=curr_row, column=9, value=r['Ty_Le_CX_%'] / 100).number_format = '0.00%'
    ws1.cell(row=curr_row, column=10, value=int(r['Thua_Lo_Gioi'])).number_format = '#,##0'
    ws1.cell(row=curr_row, column=11, value=int(r['Thua_1500'])).number_format = '#,##0'
    
    for c in range(1, 12):
        cell = ws1.cell(row=curr_row, column=c)
        cell.font = font_cell
        cell.border = border_cell
        if c > 1:
            cell.alignment = align_right
        if curr_row % 2 == 1:
            cell.fill = fill_zebra
    curr_row += 1

ws1.row_dimensions[curr_row].height = 22
ws1.cell(row=curr_row, column=1, value="TỔNG CỘNG").alignment = align_left
ws1.cell(row=curr_row, column=2, value=f"=SUM(B10:B{curr_row-1})").number_format = '#,##0'
ws1.cell(row=curr_row, column=3, value=1.0).number_format = '0.00%'
ws1.cell(row=curr_row, column=4, value=f"=SUM(D10:D{curr_row-1})").number_format = '#,##0.0'
ws1.cell(row=curr_row, column=5, value=f"=SUM(E10:E{curr_row-1})").number_format = '#,##0.00'
ws1.cell(row=curr_row, column=6, value=f"=SUM(F10:F{curr_row-1})").number_format = '#,##0'
ws1.cell(row=curr_row, column=7, value=f"=F{curr_row}/B{curr_row}").number_format = '0.00%'
ws1.cell(row=curr_row, column=8, value=f"=SUM(H10:H{curr_row-1})").number_format = '#,##0'
ws1.cell(row=curr_row, column=9, value=f"=H{curr_row}/B{curr_row}").number_format = '0.00%'
ws1.cell(row=curr_row, column=10, value=f"=SUM(J10:J{curr_row-1})").number_format = '#,##0'
ws1.cell(row=curr_row, column=11, value=f"=SUM(K10:K{curr_row-1})").number_format = '#,##0'

for c in range(1, 12):
    cell = ws1.cell(row=curr_row, column=c)
    cell.font = font_cell_bold
    cell.fill = fill_total_row
    cell.border = border_total

# SHEET 2: THỐNG KÊ 72 LOẠI ĐẤT THÔ
ws2 = wb.create_sheet(title='📚 2. Thống Kê 72 Loại Đất Thô')
ws2.views.sheetView[0].showGridLines = True

ws2.merge_cells('A1:J1')
ws2['A1'] = "BẢNG THỐNG KÊ TOÀN BỘ 72 LOẠI MỤC ĐÍCH SỬ DỤNG ĐẤT THÔ (KHÔNG GỘP - NGUYÊN BẢN SQHKT)"
ws2['A1'].font = font_title
ws2['A1'].fill = fill_navy_hdr
ws2['A1'].alignment = align_center
ws2.row_dimensions[1].height = 28

ws2.merge_cells('A2:J2')
ws2['A2'] = "Dữ liệu khảo sát thô ban đầu để xem chi tiết từng loại mục đích đất, số ô, số thửa, tổng diện tích m² và chỉ tiêu quy hoạch"
ws2['A2'].font = font_subtitle
ws2['A2'].fill = fill_navy_hdr
ws2['A2'].alignment = align_center
ws2.row_dimensions[2].height = 18

s2_headers = [
    'STT', 'Tên Loại Đất Thô (SQHKT)', 'Nhóm Tham Khảo', 'Số Ô Xuất Hiện',
    'Số Thửa Tiếp Giáp', 'Tổng Diện Tích Ô (m²)', 'Tỷ Lệ Trên Tổng Khảo Sát (%)',
    'DT Trung Bình 1 Ô (m²)', 'Mật Độ XD Max (%)', 'HSSDĐ Max'
]
ws2.row_dimensions[3].height = 24
for c_idx, h in enumerate(s2_headers, start=1):
    cell = ws2.cell(row=3, column=c_idx, value=h)
    cell.font = font_table_hdr
    cell.fill = fill_teal_hdr
    cell.alignment = align_center
    cell.border = border_header

for idx, r in raw_types_df.iterrows():
    row_num = idx + 4
    ws2.row_dimensions[row_num].height = 19
    ws2.cell(row=row_num, column=1, value=r['STT']).alignment = align_center
    ws2.cell(row=row_num, column=2, value=r['Tên Loại Đất Thô (SQHKT)']).alignment = align_left
    ws2.cell(row=row_num, column=3, value=r['Nhóm_Tham_Khảo']).alignment = align_left
    ws2.cell(row=row_num, column=4, value=int(r['Số_Ô'])).number_format = '#,##0'
    ws2.cell(row=row_num, column=5, value=int(r['Số_Thửa_Tiếp_Giáp'])).number_format = '#,##0'
    ws2.cell(row=row_num, column=6, value=r['Tổng_DT_Ô_m2']).number_format = '#,##0.0'
    ws2.cell(row=row_num, column=7, value=r['Tỷ Lệ DT Trên Tổng Khảo Sát (%)'] / 100).number_format = '0.00%'
    ws2.cell(row=row_num, column=8, value=r['DT_Trung_Bình_Ô_m2']).number_format = '#,##0.0'
    ws2.cell(row=row_num, column=9, value=r['Mật_Độ_Max'] if pd.notna(r['Mật_Độ_Max']) else '-').alignment = align_right
    ws2.cell(row=row_num, column=10, value=r['HSSDĐ_Max'] if pd.notna(r['HSSDĐ_Max']) else '-').alignment = align_right
    
    for c in range(1, 11):
        cell = ws2.cell(row=row_num, column=c)
        cell.font = font_cell
        cell.border = border_cell
        if 4 <= c <= 8:
            cell.alignment = align_right
        if row_num % 2 == 1:
            cell.fill = fill_zebra
        if 'Cây Xanh' in str(r['Nhóm_Tham_Khảo']):
            if c == 2:
                cell.fill = fill_highlight_red

ws2.auto_filter.ref = f"A3:J{len(raw_types_df)+3}"
ws2.freeze_panes = 'C4'

# SHEET 3: CHI TIẾT 10.932 Ô QH
ws3 = wb.create_sheet(title='🔍 3. Chi Tiết 10.932 Ô QH')
ws3.views.sheetView[0].showGridLines = True

ws3.merge_cells('A1:O1')
ws3['A1'] = "DANH SÁCH TOÀN BỘ 10,932 Ô PHÂN KHU CHỨC NĂNG QUY HOẠCH CHI TIẾT (ĐẦY ĐỦ NHÃN, MÃ Ô & DIỆN TÍCH)"
ws3['A1'].font = font_title
ws3['A1'].fill = fill_navy_hdr
ws3['A1'].alignment = align_center
ws3.row_dimensions[1].height = 28

s3_headers = [
    'STT Ô', 'STT Thửa', 'Số Tờ', 'Số Thửa', 'Quận / Huyện', 'Phường / Xã',
    'Mã Ô Phố', 'Chức Năng Sử Dụng Đất (Thô)', 'Chức Năng Chi Tiết (SQHKT)',
    'Nhãn Ô Chi Tiết Đầy Đủ', 'Diện Tích Ô (m²)', 'Tỷ Lệ Chiếm Thửa (%)',
    'Tầng Cao Cho Phép', 'Mật Độ XD (%)', 'HSSDĐ'
]
ws3.row_dimensions[2].height = 24
for c_idx, h in enumerate(s3_headers, start=1):
    cell = ws3.cell(row=2, column=c_idx, value=h)
    cell.font = font_table_hdr
    cell.fill = fill_blue_hdr
    cell.alignment = align_center
    cell.border = border_header

for idx, (_, r) in enumerate(df2.iterrows(), start=1):
    row_num = idx + 2
    ws3.row_dimensions[row_num].height = 18
    ws3.cell(row=row_num, column=1, value=idx).alignment = align_center
    ws3.cell(row=row_num, column=2, value=int(r['STT Thửa'])).alignment = align_center
    ws3.cell(row=row_num, column=3, value=r['Số Tờ']).alignment = align_center
    ws3.cell(row=row_num, column=4, value=r['Số Thửa']).alignment = align_center
    ws3.cell(row=row_num, column=5, value=r['Quận / Huyện']).alignment = align_left
    ws3.cell(row=row_num, column=6, value=r['Phường / Xã']).alignment = align_left
    ws3.cell(row=row_num, column=7, value=str(r['Mã Ô Phố'])).alignment = align_center
    ws3.cell(row=row_num, column=8, value=r['Chức Năng Sử Dụng Đất']).alignment = align_left
    ws3.cell(row=row_num, column=9, value=r['Chức Năng Chi Tiết (SQHKT)']).alignment = align_left
    ws3.cell(row=row_num, column=10, value=r['Nhãn Ô Chi Tiết']).alignment = align_left
    ws3.cell(row=row_num, column=11, value=r['Diện Tích Ô (m²)']).number_format = '#,##0.0'
    ws3.cell(row=row_num, column=12, value=r['Tỷ Lệ Chiếm Thửa (%)']).number_format = '0.0%'
    ws3.cell(row=row_num, column=13, value=r['Tầng Cao Cho Phép (tầng)'] if pd.notna(r['Tầng Cao Cho Phép (tầng)']) else '-').alignment = align_center
    ws3.cell(row=row_num, column=14, value=r['Mật Độ XD Tối Đa (%)'] if pd.notna(r['Mật Độ XD Tối Đa (%)']) else '-').alignment = align_right
    ws3.cell(row=row_num, column=15, value=r['Hệ Số Sử Dụng Đất (HSSDĐ)'] if pd.notna(r['Hệ Số Sử Dụng Đất (HSSDĐ)']) else '-').alignment = align_right
    
    for c in range(1, 16):
        cell = ws3.cell(row=row_num, column=c)
        cell.font = font_cell
        cell.border = border_cell
        if c in [11, 12]:
            cell.alignment = align_right
        if row_num % 2 == 1:
            cell.fill = fill_zebra
        if any(k in str(r['Chức Năng Sử Dụng Đất']).lower() for k in ['cây xanh', 'công viên', 'cv -', 'thể dục thể thao']):
            if c in [8, 9, 10]:
                cell.fill = fill_highlight_red

ws3.auto_filter.ref = f"A2:O{len(df2)+2}"
ws3.freeze_panes = 'E3'

# SHEET 4: TỔNG HỢP THỬA & CƠ CẤU
ws4 = wb.create_sheet(title='🏡 4. Tổng Hợp Thửa & Cơ Cấu')
ws4.views.sheetView[0].showGridLines = True

ws4.merge_cells('A1:R1')
ws4['A1'] = "DANH SÁCH TỔNG HỢP 6,321 THỬA ĐẤT VÀ CƠ CẤU MỤC ĐÍCH SỬ DỤNG ĐẤT CHI TIẾT"
ws4['A1'].font = font_title
ws4['A1'].fill = fill_navy_hdr
ws4['A1'].alignment = align_center
ws4.row_dimensions[1].height = 28

s4_headers = [
    'STT Thửa', 'Số Tờ', 'Số Thửa', 'Mã Thửa Đất', 'Quận / Huyện', 'Phường / Xã',
    'Tổng DT Thửa (m²)', 'Cơ Cấu Quy Hoạch Thô (SQHKT)', 'Đánh Giá An Toàn',
    'DT Đất Ở (m²)', 'Tỷ Lệ Đất Ở (%)', 'DT Cây Xanh (m²)', 'Tỷ Lệ Cây Xanh (%)',
    'DT Giao Thông (m²)', 'Tỷ Lệ Giao Thông (%)', 'DT Phức Hợp (m²)', 'Dự Án 1/500 & ĐCCB', 'Lộ Giới Tuyến Đường'
]
ws4.row_dimensions[2].height = 24
for c_idx, h in enumerate(s4_headers, start=1):
    cell = ws4.cell(row=2, column=c_idx, value=h)
    cell.font = font_table_hdr
    cell.fill = fill_blue_hdr
    cell.alignment = align_center
    cell.border = border_header

for idx, (_, r) in enumerate(df1_full_enriched.iterrows(), start=1):
    row_num = idx + 2
    ws4.row_dimensions[row_num].height = 18
    ws4.cell(row=row_num, column=1, value=int(r['STT'])).alignment = align_center
    ws4.cell(row=row_num, column=2, value=r['Số Tờ']).alignment = align_center
    ws4.cell(row=row_num, column=3, value=r['Số Thửa']).alignment = align_center
    ws4.cell(row=row_num, column=4, value=str(r['Mã Thửa Đất'])).alignment = align_left
    ws4.cell(row=row_num, column=5, value=r['Quận / Huyện']).alignment = align_left
    ws4.cell(row=row_num, column=6, value=r['Phường / Xã']).alignment = align_left
    ws4.cell(row=row_num, column=7, value=r['Diện Tích Thửa (m²)']).number_format = '#,##0.0'
    ws4.cell(row=row_num, column=8, value=r['Cơ Cấu Quy Hoạch Sử Dụng Đất']).alignment = align_left
    ws4.cell(row=row_num, column=9, value=r['Đánh Giá An Toàn']).alignment = align_left
    ws4.cell(row=row_num, column=10, value=r['DT Đất Ở (m²)']).number_format = '#,##0.0'
    ws4.cell(row=row_num, column=11, value=r['Tỷ Lệ Đất Ở (%)'] / 100).number_format = '0.0%'
    ws4.cell(row=row_num, column=12, value=r['DT Cây Xanh (m²)']).number_format = '#,##0.0'
    ws4.cell(row=row_num, column=13, value=r['Tỷ Lệ Cây Xanh (%)'] / 100).number_format = '0.0%'
    ws4.cell(row=row_num, column=14, value=r['DT Giao Thông Lộ Giới (m²)']).number_format = '#,##0.0'
    ws4.cell(row=row_num, column=15, value=r['Tỷ Lệ Giao Thông (%)'] / 100).number_format = '0.0%'
    ws4.cell(row=row_num, column=16, value=r['DT Phức Hợp (m²)']).number_format = '#,##0.0'
    ws4.cell(row=row_num, column=17, value=r['Dự Án 1/500 & ĐCCB']).alignment = align_center
    ws4.cell(row=row_num, column=18, value=r['Lộ Giới Tuyến Đường Tiếp Giáp']).alignment = align_left
    
    for c in range(1, 19):
        cell = ws4.cell(row=row_num, column=c)
        cell.font = font_cell
        cell.border = border_cell
        if c in [7, 10, 11, 12, 13, 14, 15, 16]:
            cell.alignment = align_right
        if row_num % 2 == 1:
            cell.fill = fill_zebra
        if r['DT Cây Xanh (m²)'] > 0 and c in [12, 13]:
            cell.fill = fill_highlight_red
            cell.font = font_cell_bold

ws4.auto_filter.ref = f"A2:R{len(df1_full_enriched)+2}"
ws4.freeze_panes = 'E3'

# SHEET 5: THỬA KO DÍNH CÂY XANH
ws5 = wb.create_sheet(title='🌿 5. Thửa KO Dính Cây Xanh')
ws5.views.sheetView[0].showGridLines = True

df_clean = df1_full_enriched[df1_full_enriched['DT Cây Xanh (m²)'] == 0].copy()

ws5.merge_cells('A1:N1')
ws5['A1'] = "DANH SÁCH 6,139 THỬA ĐẤT AN TOÀN - HOÀN TOÀN KHÔNG DÍNH QUY HOẠCH CÂY XANH / CÔNG VIÊN"
ws5['A1'].font = font_title
ws5['A1'].fill = fill_navy_hdr
ws5['A1'].alignment = align_center
ws5.row_dimensions[1].height = 28

s5_headers = [
    'STT Thửa', 'Số Tờ', 'Số Thửa', 'Mã Thửa Đất', 'Quận / Huyện', 'Phường / Xã',
    'Tổng DT Thửa (m²)', 'Đánh Giá An Toàn', 'DT Đất Ở (m²)', 'Tỷ Lệ Đất Ở (%)',
    'DT Giao Thông (m²)', 'DT Phức Hợp (m²)', 'Dự Án 1/500 & ĐCCB', 'Cơ Cấu Quy Hoạch Thô'
]
ws5.row_dimensions[2].height = 24
for c_idx, h in enumerate(s5_headers, start=1):
    cell = ws5.cell(row=2, column=c_idx, value=h)
    cell.font = font_table_hdr
    cell.fill = fill_teal_hdr
    cell.alignment = align_center
    cell.border = border_header

for idx, (_, r) in enumerate(df_clean.iterrows(), start=1):
    row_num = idx + 2
    ws5.row_dimensions[row_num].height = 18
    ws5.cell(row=row_num, column=1, value=int(r['STT'])).alignment = align_center
    ws5.cell(row=row_num, column=2, value=r['Số Tờ']).alignment = align_center
    ws5.cell(row=row_num, column=3, value=r['Số Thửa']).alignment = align_center
    ws5.cell(row=row_num, column=4, value=str(r['Mã Thửa Đất'])).alignment = align_left
    ws5.cell(row=row_num, column=5, value=r['Quận / Huyện']).alignment = align_left
    ws5.cell(row=row_num, column=6, value=r['Phường / Xã']).alignment = align_left
    ws5.cell(row=row_num, column=7, value=r['Diện Tích Thửa (m²)']).number_format = '#,##0.0'
    ws5.cell(row=row_num, column=8, value=r['Đánh Giá An Toàn']).alignment = align_left
    ws5.cell(row=row_num, column=9, value=r['DT Đất Ở (m²)']).number_format = '#,##0.0'
    ws5.cell(row=row_num, column=10, value=r['Tỷ Lệ Đất Ở (%)'] / 100).number_format = '0.0%'
    ws5.cell(row=row_num, column=11, value=r['DT Giao Thông Lộ Giới (m²)']).number_format = '#,##0.0'
    ws5.cell(row=row_num, column=12, value=r['DT Phức Hợp (m²)']).number_format = '#,##0.0'
    ws5.cell(row=row_num, column=13, value=r['Dự Án 1/500 & ĐCCB']).alignment = align_center
    ws5.cell(row=row_num, column=14, value=r['Cơ Cấu Quy Hoạch Sử Dụng Đất']).alignment = align_left
    
    for c in range(1, 15):
        cell = ws5.cell(row=row_num, column=c)
        cell.font = font_cell
        cell.border = border_cell
        if c in [7, 9, 10, 11, 12]:
            cell.alignment = align_right
        if row_num % 2 == 1:
            cell.fill = fill_zebra

ws5.auto_filter.ref = f"A2:N{len(df_clean)+2}"
ws5.freeze_panes = 'E3'

# SHEET 6: THỬA DÍNH CÂY XANH
ws6 = wb.create_sheet(title='⚠️ 6. Thửa Dính Cây Xanh')
ws6.views.sheetView[0].showGridLines = True

df_green_parcels = df1_full_enriched[df1_full_enriched['DT Cây Xanh (m²)'] > 0].sort_values(by='Tỷ Lệ Cây Xanh (%)', ascending=False).copy()

ws6.merge_cells('A1:M1')
ws6['A1'] = "DANH SÁCH 182 THỬA ĐẤT BỊ DÍNH QUY HOẠCH CÂY XANH, CÔNG VIÊN, CÁCH LY (CẢNH BÁO RỦI RO)"
ws6['A1'].font = font_title
ws6['A1'].fill = PatternFill(start_color='991B1B', end_color='991B1B', fill_type='solid')
ws6['A1'].alignment = align_center
ws6.row_dimensions[1].height = 28

s6_headers = [
    'STT', 'Mã Thửa Đất', 'Số Tờ', 'Số Thửa', 'Quận / Huyện', 'Phường / Xã',
    'Tổng DT Thửa (m²)', 'Loại Cây Xanh Thô', 'Chi Tiết Ô Cây Xanh Dính Phải',
    'DT Cây Xanh (m²)', 'Tỷ Lệ Cây Xanh (%)', 'DT Đất Ở Còn Lại (m²)', 'Tỷ Lệ Đất Ở (%)'
]
ws6.row_dimensions[2].height = 24
for c_idx, h in enumerate(s6_headers, start=1):
    cell = ws6.cell(row=2, column=c_idx, value=h)
    cell.font = font_table_hdr
    cell.fill = PatternFill(start_color='B91C1C', end_color='B91C1C', fill_type='solid')
    cell.alignment = align_center
    cell.border = border_header

for idx, (_, r) in enumerate(df_green_parcels.iterrows(), start=1):
    row_num = idx + 2
    ws6.row_dimensions[row_num].height = 18
    ws6.cell(row=row_num, column=1, value=idx).alignment = align_center
    ws6.cell(row=row_num, column=2, value=str(r['Mã Thửa Đất'])).alignment = align_left
    ws6.cell(row=row_num, column=3, value=r['Số Tờ']).alignment = align_center
    ws6.cell(row=row_num, column=4, value=r['Số Thửa']).alignment = align_center
    ws6.cell(row=row_num, column=5, value=r['Quận / Huyện']).alignment = align_left
    ws6.cell(row=row_num, column=6, value=r['Phường / Xã']).alignment = align_left
    ws6.cell(row=row_num, column=7, value=r['Diện Tích Thửa (m²)']).number_format = '#,##0.0'
    ws6.cell(row=row_num, column=8, value=r['Loại Cây Xanh Thô']).alignment = align_left
    ws6.cell(row=row_num, column=9, value=r['Chi Tiết Ô Cây Xanh']).alignment = align_left
    ws6.cell(row=row_num, column=10, value=r['DT Cây Xanh (m²)']).number_format = '#,##0.0'
    ws6.cell(row=row_num, column=11, value=r['Tỷ Lệ Cây Xanh (%)'] / 100).number_format = '0.0%'
    ws6.cell(row=row_num, column=12, value=r['DT Đất Ở (m²)']).number_format = '#,##0.0'
    ws6.cell(row=row_num, column=13, value=r['Tỷ Lệ Đất Ở (%)'] / 100).number_format = '0.0%'
    
    for c in range(1, 14):
        cell = ws6.cell(row=row_num, column=c)
        cell.font = font_cell
        cell.border = border_cell
        if c in [7, 10, 11, 12, 13]:
            cell.alignment = align_right
        if row_num % 2 == 1:
            cell.fill = fill_highlight_red

ws6.auto_filter.ref = f"A2:M{len(df_green_parcels)+2}"
ws6.freeze_panes = 'E3'

# SHEET 7: DỰ ÁN 1-500 & ĐCCB
ws7 = wb.create_sheet(title='🏗️ 7. Dự Án 1-500 & DCCB')
ws7.views.sheetView[0].showGridLines = True

ws7.merge_cells('A1:L1')
ws7['A1'] = "DANH MỤC 632 BẢN GHI THUỘC ĐỒ ÁN QUY HOẠCH 1/500 & ĐIỀU CHỈNH CỤC BỘ"
ws7['A1'].font = font_title
ws7['A1'].fill = fill_navy_hdr
ws7['A1'].alignment = align_center
ws7.row_dimensions[1].height = 28

s7_headers = [
    'STT', 'Mã Thửa Đất', 'Số Tờ', 'Số Thửa', 'Quận / Huyện', 'Phường / Xã',
    'Loại Quy Hoạch', 'Tên Dự Án / Quyết Định Điều Chỉnh', 'Số Quyết Định',
    'Ngày Phê Duyệt', 'Cơ Quan Phê Duyệt', 'Tỷ Lệ Thửa Thuộc Dự Án (%)'
]
ws7.row_dimensions[2].height = 24
for c_idx, h in enumerate(s7_headers, start=1):
    cell = ws7.cell(row=2, column=c_idx, value=h)
    cell.font = font_table_hdr
    cell.fill = fill_purple_hdr
    cell.alignment = align_center
    cell.border = border_header

for idx, (_, r) in enumerate(df4.iterrows(), start=1):
    row_num = idx + 2
    ws7.row_dimensions[row_num].height = 18
    ws7.cell(row=row_num, column=1, value=idx).alignment = align_center
    ws7.cell(row=row_num, column=2, value=str(r['Mã Thửa Đất'])).alignment = align_left
    ws7.cell(row=row_num, column=3, value=r['Số Tờ']).alignment = align_center
    ws7.cell(row=row_num, column=4, value=r['Số Thửa']).alignment = align_center
    ws7.cell(row=row_num, column=5, value=r['Quận / Huyện']).alignment = align_left
    ws7.cell(row=row_num, column=6, value=r['Phường / Xã']).alignment = align_left
    ws7.cell(row=row_num, column=7, value=r['Loại Quy Hoạch']).alignment = align_center
    ws7.cell(row=row_num, column=8, value=r['Tên Dự Án / Đồ Án Điều Chỉnh']).alignment = align_left
    ws7.cell(row=row_num, column=9, value=r['Số Quyết Định Phê Duyệt'] if pd.notna(r['Số Quyết Định Phê Duyệt']) else '-').alignment = align_left
    ws7.cell(row=row_num, column=10, value=str(r['Ngày Phê Duyệt']) if pd.notna(r['Ngày Phê Duyệt']) else '-').alignment = align_center
    ws7.cell(row=row_num, column=11, value=r['Cơ Quan Phê Duyệt'] if pd.notna(r['Cơ Quan Phê Duyệt']) else '-').alignment = align_left
    
    pct_val = r['Tỷ Lệ Thửa Thuộc Dự Án (%)']
    if pd.notna(pct_val):
        ws7.cell(row=row_num, column=12, value=float(pct_val) if float(pct_val) <= 1.0 else float(pct_val)/100).number_format = '0.0%'
    else:
        ws7.cell(row=row_num, column=12, value='-').alignment = align_center
    
    for c in range(1, 13):
        cell = ws7.cell(row=row_num, column=c)
        cell.font = font_cell
        cell.border = border_cell
        if c == 12:
            cell.alignment = align_right
        if row_num % 2 == 1:
            cell.fill = fill_zebra
        if '1/500' in str(r['Loại Quy Hoạch']):
            if c == 7:
                cell.fill = fill_highlight_amber

ws7.auto_filter.ref = f"A2:L{len(df4)+2}"
ws7.freeze_panes = 'E3'

# Auto widths
print("8. Auto-adjusting column widths...")
for sheet in wb.worksheets:
    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in [1, 2] and col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q', 'R']:
                continue
            if cell.value:
                val_str = str(cell.value)
                max_len = max(max_len, len(val_str))
        sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 65)

ws1.column_dimensions['A'].width = 30
ws1.column_dimensions['B'].width = 18
ws1.column_dimensions['C'].width = 18
ws1.column_dimensions['D'].width = 22
ws1.column_dimensions['E'].width = 16
ws1.column_dimensions['F'].width = 18
ws1.column_dimensions['G'].width = 38

print(f"9. Saving workbook to: {output_excel}")
wb.save(output_excel)
print("SUCCESS: Clean Excel report with unmerged granular categories created successfully!")
